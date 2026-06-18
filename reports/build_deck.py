# -*- coding: utf-8 -*-
"""
Understand-Anything 深度洞察报告 — PPTX 生成脚本
Egonex-AI / Understand-Anything 开源项目分析
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ----------------------------------------------------------------------------
# Palette
# ----------------------------------------------------------------------------
BG     = RGBColor(0x0E, 0x17, 0x26)   # deep navy
BG2    = RGBColor(0x15, 0x20, 0x33)   # panel
BG3    = RGBColor(0x1C, 0x2A, 0x42)   # panel lighter
INK    = RGBColor(0xEA, 0xF0, 0xF7)   # near white
MUTED  = RGBColor(0x96, 0xA6, 0xBE)   # muted gray-blue
FAINT  = RGBColor(0x5E, 0x6E, 0x86)
TEAL   = RGBColor(0x2D, 0xD4, 0xBF)
CYAN   = RGBColor(0x38, 0xBD, 0xF8)
AMBER  = RGBColor(0xFB, 0xBF, 0x24)
VIOLET = RGBColor(0xA7, 0x8B, 0xFA)
RED    = RGBColor(0xF8, 0x71, 0x71)
GREEN  = RGBColor(0x34, 0xD3, 0x99)
PINK   = RGBColor(0xF4, 0x72, 0xB6)

LAYER = {"API": CYAN, "Service": TEAL, "Data": AMBER, "UI": VIOLET, "Utility": GREEN}

LATIN = "Calibri"
EA    = "Microsoft YaHei"   # CJK-capable, common on Windows; falls back gracefully

EMUIN = 914400
SW, SH = 13.333, 7.5  # widescreen inches

prs = Presentation()
prs.slide_width  = Emu(int(SW * EMUIN))
prs.slide_height = Emu(int(SH * EMUIN))
BLANK = prs.slide_layouts[6]

# ----------------------------------------------------------------------------
# Low-level helpers
# ----------------------------------------------------------------------------
def _fonts(run, latin=LATIN, ea=EA):
    run.font.name = latin
    rPr = run._r.get_or_add_rPr()
    for tag, face in (('a:latin', latin), ('a:ea', ea), ('a:cs', latin)):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set('typeface', face)

def slide():
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = BG
    r.line.fill.background(); r.shadow.inherit = False
    return s

def rect(s, l, t, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE, radius=None):
    shp = s.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try: shp.adjustments[0] = radius
        except Exception: pass
    return shp

def txt(s, l, t, w, h, lines, anchor=MSO_ANCHOR.TOP):
    """lines: list of dicts {text,size,color,bold,align,bullet,sb,sa,level,ea,latin}"""
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ln.get('align', PP_ALIGN.LEFT)
        p.level = ln.get('level', 0)
        if 'sb' in ln: p.space_before = Pt(ln['sb'])
        if 'sa' in ln: p.space_after = Pt(ln['sa'])
        if 'lh' in ln: p.line_spacing = ln['lh']
        segs = ln.get('segs')
        if segs is None:
            t0 = ln.get('text', '')
            if ln.get('bullet'): t0 = ln['bullet'] + '  ' + t0
            segs = [(t0, ln)]
        for text, st in segs:
            r = p.add_run(); r.text = text
            r.font.size = Pt(st.get('size', 14))
            r.font.bold = st.get('bold', False)
            r.font.italic = st.get('italic', False)
            r.font.color.rgb = st.get('color', INK)
            _fonts(r, st.get('latin', LATIN), st.get('ea', EA))
    return tb

def fill_shape_text(shp, lines, anchor=MSO_ANCHOR.MIDDLE):
    tf = shp.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.08); tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04); tf.margin_bottom = Inches(0.04)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ln.get('align', PP_ALIGN.CENTER)
        if 'sb' in ln: p.space_before = Pt(ln['sb'])
        if 'sa' in ln: p.space_after = Pt(ln['sa'])
        if 'lh' in ln: p.line_spacing = ln['lh']
        r = p.add_run(); r.text = ln.get('text', '')
        r.font.size = Pt(ln.get('size', 12)); r.font.bold = ln.get('bold', False)
        r.font.color.rgb = ln.get('color', INK)
        _fonts(r, ln.get('latin', LATIN), ln.get('ea', EA))

def header(s, kicker, title, idx):
    rect(s, 0.55, 0.52, 0.10, 0.62, fill=TEAL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    txt(s, 0.78, 0.46, 11.6, 0.42, [{'text': kicker, 'size': 11.5, 'color': TEAL, 'bold': True}])
    txt(s, 0.78, 0.70, 11.8, 0.62, [{'text': title, 'size': 25, 'color': INK, 'bold': True}])
    footer(s, idx)

def footer(s, idx):
    rect(s, 0.55, 7.06, 12.23, 0.012, fill=BG3)
    txt(s, 0.55, 7.10, 9.0, 0.3, [{'text': 'Understand-Anything 深度洞察报告  ·  Egonex-AI', 'size': 8.5, 'color': FAINT}])
    txt(s, 11.0, 7.10, 1.78, 0.3, [{'text': f'{idx:02d} / 17   ·   2026-06-18', 'size': 8.5, 'color': FAINT, 'align': PP_ALIGN.RIGHT}])

def chip(s, l, t, w, h, label, color, fontsize=10.5, fill=None):
    c = rect(s, l, t, w, h, fill=fill if fill else BG3, line=color, lw=1.0,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    fill_shape_text(c, [{'text': label, 'size': fontsize, 'color': color, 'bold': True}])
    return c

def styled_table(s, l, t, w, rows, col_w, header_fill=BG3, header_color=TEAL,
                 body_color=INK, fontsize=10.5, row_h=0.34, head_h=0.40):
    nrows, ncols = len(rows), len(rows[0])
    gtbl = s.shapes.add_table(nrows, ncols, Inches(l), Inches(t), Inches(w), Inches(head_h + row_h*(nrows-1)))
    tbl = gtbl.table
    # disable banding style
    tbl.first_row = False; tbl.horz_banding = False
    for ci, cw in enumerate(col_w):
        tbl.columns[ci].width = Inches(cw)
    tbl.rows[0].height = Inches(head_h)
    for ri in range(1, nrows):
        tbl.rows[ri].height = Inches(row_h)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if ri == 0 else (BG2 if ri % 2 else RGBColor(0x12,0x1C,0x2E))
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.08); cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.02); cell.margin_bottom = Inches(0.02)
            tf = cell.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
            r = p.add_run(); r.text = str(val)
            r.font.size = Pt(fontsize if ri else fontsize+0.5)
            r.font.bold = (ri == 0)
            r.font.color.rgb = header_color if ri == 0 else body_color
            _fonts(r)
    return tbl

# ============================================================================
# SLIDE 1 — Cover
# ============================================================================
s = slide()
# decorative knowledge-graph on the right
import math
nodes = [(10.0,1.5,TEAL,0.16),(11.3,2.3,CYAN,0.12),(9.4,2.9,VIOLET,0.12),(11.9,3.6,AMBER,0.10),
         (10.6,4.0,TEAL,0.13),(12.3,1.7,GREEN,0.10),(9.0,4.6,CYAN,0.11),(11.0,5.2,VIOLET,0.10),
         (12.5,4.7,TEAL,0.10),(10.1,5.9,AMBER,0.09)]
edges = [(0,1),(0,2),(1,3),(0,4),(1,5),(2,6),(4,7),(3,8),(4,8),(6,7),(7,9),(2,4)]
for a,b in edges:
    x1,y1 = nodes[a][0]+nodes[a][3]/2, nodes[a][1]+nodes[a][3]/2
    x2,y2 = nodes[b][0]+nodes[b][3]/2, nodes[b][1]+nodes[b][3]/2
    cn = s.shapes.add_connector(2, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    cn.line.color.rgb = BG3; cn.line.width = Pt(1.4); cn.shadow.inherit = False
for x,y,c,d in nodes:
    rect(s, x, y, d, d, fill=c, shape=MSO_SHAPE.OVAL)

rect(s, 0.85, 1.55, 0.10, 1.7, fill=TEAL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
txt(s, 1.15, 1.55, 8.2, 0.5, [{'text': '开源项目 · 深度洞察报告', 'size': 14, 'color': TEAL, 'bold': True}])
txt(s, 1.13, 2.05, 8.6, 1.9, [
    {'text': 'Understand-Anything', 'size': 49, 'color': INK, 'bold': True},
    {'text': '把任意代码库变成可探索、可搜索、可问答的知识图谱', 'size': 17, 'color': MUTED, 'sb': 10},
])
txt(s, 1.15, 4.35, 8.4, 0.9, [
    {'segs': [('Egonex-AI / Understand-Anything', {'size':13,'color':CYAN,'bold':True}),
              ('   ·   TypeScript Monorepo   ·   MIT', {'size':13,'color':MUTED})]},
    {'segs': [('Tree-sitter + LLM 混合引擎  ·  9-Agent 流水线  ·  15+ AI 平台插件', {'size':12,'color':MUTED})], 'sb': 6},
])
# metric strip
mx = 1.15
for label, val, col in [('GitHub Stars','≈ 62.9k ★', AMBER), ('创建于','2026-03-15', TEAL),
                        ('当前版本','v2.8.0', CYAN), ('License','MIT', GREEN)]:
    b = rect(s, mx, 5.55, 1.95, 0.92, fill=BG2, line=BG3, lw=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    fill_shape_text(b, [{'text': val,'size':15,'color':col,'bold':True},
                        {'text': label,'size':9.5,'color':MUTED,'sb':2}])
    mx += 2.07
txt(s, 1.15, 6.75, 9.0, 0.3, [{'text': '分析日期 2026-06-18   ·   数据经 GitHub API 校验   ·   workflow 多 Agent 调研模式', 'size': 9, 'color': FAINT}])

# ============================================================================
# SLIDE 2 — Executive Summary
# ============================================================================
s = slide(); header(s, 'EXECUTIVE SUMMARY', '执行摘要:一句话看懂这个项目', 2)
# left big statement
p = rect(s, 0.55, 1.62, 5.5, 4.0, fill=BG2, line=BG3, lw=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
txt(s, 0.85, 1.85, 4.95, 3.6, [
    {'text': '它是什么', 'size': 12, 'color': TEAL, 'bold': True},
    {'text': '一个本地运行的 AI 代码理解工具。用确定性静态解析(Tree-sitter)抽取代码结构,再用 LLM 补充语义,经多 Agent 流水线产出一份可提交到 Git 的知识图谱 JSON,并配交互式可视化看板。', 'size': 13, 'color': INK, 'sb': 6, 'lh': 1.18},
    {'text': '解决什么痛点', 'size': 12, 'color': CYAN, 'bold': True, 'sb': 14},
    {'text': '大型/陌生代码库的「上手难」与「改动影响难评估」——既服务于人(新人 onboarding),也服务于 AI Agent(用图谱降低检索 token 与工具调用)。', 'size': 13, 'color': INK, 'sb': 6, 'lh': 1.18},
])
# right: verdict + facts
txt(s, 6.35, 1.60, 6.45, 0.4, [{'text': '分析师判断', 'size': 12, 'color': AMBER, 'bold': True}])
for i,(t,c) in enumerate([
    ('技术架构扎实且有真实差异化:混合解析 + 可提交图谱 + 多平台分发是站得住的设计。', GREEN),
    ('生态打法是最强护城河:覆盖 15+ AI 编码平台,广度超过多数同类。', GREEN),
    ('但厂商极新、身份单薄(组织仅 1 仓库),3 个月冲到 ~63k★ 的增速应作「热度信号」而非「采纳证据」看待。', AMBER),
    ('结论:适合先行试用于 onboarding / Agent 上下文增强;暂不建议作为企业关键路径的唯一依赖。', CYAN),
]):
    b = rect(s, 6.35, 2.02 + i*0.84, 6.45, 0.74, fill=BG2, line=c, lw=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    txt(s, 6.55, 2.10 + i*0.84, 6.10, 0.62, [{'text': t, 'size': 11.5, 'color': INK, 'lh': 1.08}], anchor=MSO_ANCHOR.MIDDLE)
txt(s, 0.85, 5.78, 11.9, 0.9, [
    {'segs': [('一句话总结:  ', {'size':12.5,'color':TEAL,'bold':True}),
              ('"graphs that teach > graphs that impress" —— 它把代码理解从「读代码」升级为「探索一张会教你的图」,并让这张图成为团队可共享的资产。', {'size':12.5,'color':INK})], 'lh':1.15},
])

# ============================================================================
# SLIDE 3 — Project Snapshot
# ============================================================================
s = slide(); header(s, 'PROJECT SNAPSHOT', '项目快照:关键事实与数据', 3)
rows = [
    ['维度', '数据 / 事实'],
    ['仓库', 'Egonex-AI/Understand-Anything  ·  homepage: understand-anything.com'],
    ['Stars / Forks', '≈ 62,928 ★  /  5,188 fork   (GitHub API 校验)'],
    ['创建 / 最近推送', '2026-03-15  /  2026-06-18   —— 约 3 个月,活跃维护中'],
    ['Open Issues', '212'],
    ['License / 版本', 'MIT  ·  插件 v2.8.0  ·  @understand-anything/core v0.1.0'],
    ['语言构成', 'TypeScript ~71%  ·  JavaScript ~16%  ·  Python ~9%  ·  Astro ~2%'],
    ['工程形态', 'pnpm 多包 Monorepo(core / dashboard / tree-sitter-dart-wasm / skills…)'],
    ['厂商', 'Egonex(组织仅 1 公开仓库、~48 followers;Twitter @EgonexAI;affiliate@egonex.ai)'],
]
styled_table(s, 0.55, 1.60, 7.45, rows, [1.55, 5.90], fontsize=10.2, row_h=0.46, head_h=0.40)

# caution panel
cb = rect(s, 8.25, 1.60, 4.55, 2.55, fill=RGBColor(0x2A,0x20,0x12), line=AMBER, lw=1.4, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
txt(s, 8.50, 1.78, 4.05, 2.3, [
    {'text': '⚠  星标增速需谨慎解读', 'size': 12.5, 'color': AMBER, 'bold': True},
    {'text': '3 个月累积 ~63k★ 极不寻常。数据真实(API 可查),但在「代码知识图谱」这一 2026 爆火赛道,此类增速曾引发行业质疑。', 'size': 11, 'color': INK, 'sb': 8, 'lh': 1.18},
    {'text': '建议把它当「可见度 / 热度」信号,而非已验证的生产采纳指标。贡献者深度无法独立确认。', 'size': 11, 'color': MUTED, 'sb': 8, 'lh': 1.18},
])
# verified vs unconfirmed
vb = rect(s, 8.25, 4.32, 4.55, 2.55, fill=BG2, line=BG3, lw=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
txt(s, 8.50, 4.48, 4.05, 2.3, [
    {'segs':[('✓ 已校验  ', {'size':11.5,'color':GREEN,'bold':True}), ('(GitHub API)', {'size':10,'color':MUTED})]},
    {'text': 'stars / forks / license / 创建与推送日期 / 语言构成', 'size': 10.5, 'color': INK, 'sb': 3, 'lh': 1.15},
    {'segs':[('◐ 二手来源  ', {'size':11.5,'color':AMBER,'bold':True}), ('(网络资料)', {'size':10,'color':MUTED})], 'sb': 9},
    {'text': '发布版本数、竞品星数、市场规模与 onboarding 时长等行业数据', 'size': 10.5, 'color': INK, 'sb': 3, 'lh': 1.15},
    {'segs':[('✗ 未披露  ', {'size':11.5,'color':RED,'bold':True})], 'sb': 9},
    {'text': '团队成员 / 公司主体 / 融资情况 / 贡献者规模', 'size': 10.5, 'color': INK, 'sb': 3, 'lh': 1.15},
])

# ============================================================================
# SLIDE 4 — Problem & Positioning
# ============================================================================
s = slide(); header(s, 'PROBLEM & POSITIONING', '要解决的问题:两类买家,一个痛点', 4)
txt(s, 0.55, 1.55, 12.2, 0.5, [{'text': '核心痛点:大型/陌生代码库——读不完、关系看不清、改动影响难评估。该工具同时服务于「人」与「AI Agent」两类买家。', 'size': 12.5, 'color': MUTED}])
# two buyer cards
cards = [
    ('① 人:onboarding 与影响分析', CYAN, [
        '新人/跨团队成员快速建立架构心智,「数周 → 数小时」',
        'PR / 改动的波及范围(ripple / blast-radius)评估',
        '依赖顺序的「导览路径」+ 业务域视图,边看边学',
        '买家:工程负责人、平台/DevEx 团队、开源维护者',
    ]),
    ('② AI Agent:高效结构化上下文', TEAL, [
        '用图谱回答 grep/embedding 难答的结构问题:谁调用了 X、Y 的影响面',
        '显著降低 Agent 的 token 与工具调用开销(2026 头号卖点)',
        '本地优先、无代码外泄,省去 embedding API 成本',
        '买家:大规模使用 Claude Code/Cursor/Codex 的团队',
    ]),
]
for i,(title, col, items) in enumerate(cards):
    l = 0.55 + i*6.25
    c = rect(s, l, 2.20, 5.95, 4.05, fill=BG2, line=col, lw=1.3, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    rect(s, l, 2.20, 5.95, 0.62, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
    txt(s, l+0.25, 2.30, 5.5, 0.45, [{'text': title, 'size': 14, 'color': BG, 'bold': True}])
    lines = []
    for it in items:
        lines.append({'text': it, 'size': 12, 'color': INK, 'bullet': '▸', 'sb': 9, 'lh': 1.12})
    txt(s, l+0.30, 3.05, 5.4, 3.05, lines)
txt(s, 0.55, 6.45, 12.2, 0.5, [{'segs':[
    ('定位精髓: ', {'size':12,'color':AMBER,'bold':True}),
    ('它同时站在「Agent 上下文供给」与「人类代码理解/onboarding」两个相邻赛道的交叉点——这是它区别于纯检索工具与纯可视化工具的关键。', {'size':12,'color':INK})]}])

# ============================================================================
# SLIDE 5 — Design Philosophy
# ============================================================================
s = slide(); header(s, 'DESIGN PHILOSOPHY', '设计哲学:确定性骨架 + LLM 血肉', 5)
# two columns: tree-sitter vs LLM
colL = rect(s, 0.55, 1.65, 5.95, 3.05, fill=BG2, line=TEAL, lw=1.3, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
txt(s, 0.80, 1.82, 5.5, 2.8, [
    {'text': '结构事实  ·  Tree-sitter 静态解析', 'size': 13.5, 'color': TEAL, 'bold': True},
    {'text': 'imports / 函数 / 类 / 调用点 / 端点', 'size': 11.5, 'color': MUTED, 'sb': 4},
    {'text': '保证:same input → same output(可复现)', 'size': 12, 'color': INK, 'bullet':'▸', 'sb': 12, 'lh':1.12},
    {'text': '无幻觉:路径/定义来自真实 AST,Agent 被严令「绝不编造文件路径」', 'size': 12, 'color': INK, 'bullet':'▸', 'sb': 8, 'lh':1.12},
    {'text': '由 .mjs 脚本确定性执行,输出可逐字节比对', 'size': 12, 'color': INK, 'bullet':'▸', 'sb': 8, 'lh':1.12},
])
colR = rect(s, 6.85, 1.65, 5.95, 3.05, fill=BG2, line=VIOLET, lw=1.3, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
txt(s, 7.10, 1.82, 5.5, 2.8, [
    {'text': '语义意图  ·  LLM 理解', 'size': 13.5, 'color': VIOLET, 'bold': True},
    {'text': '摘要 / 架构层归属 / 业务域映射 / 复杂度', 'size': 11.5, 'color': MUTED, 'sb': 4},
    {'text': '回答「为什么」:每个节点配纯语言摘要', 'size': 12, 'color': INK, 'bullet':'▸', 'sb': 12, 'lh':1.12},
    {'text': '把零散结构组织成「会教人」的层次与导览', 'size': 12, 'color': INK, 'bullet':'▸', 'sb': 8, 'lh':1.12},
    {'text': '多 Agent 协作 + Reviewer 校验,缓解漂移', 'size': 12, 'color': INK, 'bullet':'▸', 'sb': 8, 'lh':1.12},
])
# fusion bar
rect(s, 0.55, 4.95, 12.25, 0.66, fill=BG3, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2)
txt(s, 0.55, 4.95, 12.25, 0.66, [{'segs':[
    ('融合 →  ', {'size':13,'color':AMBER,'bold':True}),
    ('可复现的结构骨架  +  富语义的解释血肉  =  既可信、又好懂的知识图谱', {'size':13,'color':INK,'bold':True})],
    'align': PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
# tenets
for i,(t,c) in enumerate([('「教育 > 炫技」的图谱设计', TEAL),('图谱即代码:可提交、可共享', CYAN),('Persona 自适应:按角色调复杂度', VIOLET)]):
    chip(s, 0.55 + i*4.13, 5.95, 3.95, 0.66, t, c, fontsize=12)

# ============================================================================
# SLIDE 6 — Architecture Overview
# ============================================================================
s = slide(); header(s, 'ARCHITECTURE', '系统架构总览:五层结构', 6)
bands = [
    ('入口层  ·  CLI / IDE 插件分发', CYAN, '原生 Claude Code 插件  +  Cursor · VS Code+Copilot · Codex · Gemini CLI · Cline · Trae · Kiro …  15+ 平台(install.sh + 符号链接自动发现)'),
    ('编排层  ·  9-Agent 流水线 + 确定性脚本', TEAL, 'scan-project / extract-structure / extract-import-map / compute-batches / build-fingerprints (.mjs)  +  merge-batch-graphs (.py) ;  Agents 并行(5 并发,20–30 文件/批)'),
    ('引擎层  ·  三引擎协同', AMBER, 'Tree-sitter(web-tree-sitter WASM,12 语言)结构解析   |   LLM 语义理解   |   graphology + Louvain 图算法/社区聚类'),
    ('数据层  ·  可提交知识图谱', VIOLET, '.understand-anything/knowledge-graph.json(version/project/nodes/edges/layers/tour)  +  intermediate/ 中间产物  ·  fingerprint & staleness 支撑增量更新'),
    ('展现层  ·  交互式 Dashboard', GREEN, 'React 19 + Vite 6 + Tailwind 4 + Zustand ;  React Flow(@xyflow/react)+ ELK 两阶段层次布局 ;  模糊+语义搜索 / 导览 / diff 叠加 / persona UI'),
]
y = 1.58
for title, col, desc in bands:
    h = 0.96
    rect(s, 0.55, y, 12.25, h, fill=BG2, line=BG3, lw=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    rect(s, 0.55, y, 0.12, h, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    txt(s, 0.85, y+0.13, 11.7, 0.4, [{'text': title, 'size': 13, 'color': col, 'bold': True}])
    txt(s, 0.85, y+0.48, 11.8, 0.45, [{'text': desc, 'size': 10.8, 'color': MUTED, 'lh': 1.08}])
    if y < 6.0:
        tri = rect(s, 6.55, y+h-0.04, 0.24, 0.16, fill=BG3, shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
        tri.rotation = 180
    y += h + 0.14

# ============================================================================
# SLIDE 7 — 9-Agent Pipeline
# ============================================================================
s = slide(); header(s, 'CORE WORKFLOW', '核心 Workflow:9-Agent 协作流水线', 7)
txt(s, 0.55, 1.52, 12.2, 0.4, [{'segs':[
    ('/understand 触发的核心分析流水线（6 个） ', {'size':12,'color':TEAL,'bold':True}),
    ('—— 每个 Agent 先跑确定性脚本拿事实,再做受约束的 LLM 合成', {'size':11,'color':MUTED})]}])
core = [
    ('1','Project-\nScanner','清点文件、识别语言/框架;严禁编造路径'),
    ('2','File-\nAnalyzer','tree-sitter 抽函数/类/调用→节点+边(分批)'),
    ('3','Assemble-\nReviewer','合并批次图、补回丢失节点/边'),
    ('4','Architecture-\nAnalyzer','按 70+ 目录约定分配 3–10 个架构层'),
    ('5','Tour-\nBuilder','按依赖生成 5–15 步导览路径'),
    ('6','Graph-\nReviewer','9 类校验,critical 阻断、空问题才通过'),
]
n = len(core); gap = 0.16; bw = (12.25 - gap*(n-1))/n; bh = 1.95; y = 2.10
rect(s, 0.55, y+bh/2-0.015, 12.25, 0.03, fill=BG3)  # pipeline rail
for i,(num,name,role) in enumerate(core):
    l = 0.55 + i*(bw+gap)
    b = rect(s, l, y, bw, bh, fill=BG2, line=TEAL, lw=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    txt(s, l, y+0.14, bw, 0.5, [{'text': num, 'size': 22, 'color': TEAL, 'bold': True, 'align': PP_ALIGN.CENTER}])
    txt(s, l+0.06, y+0.66, bw-0.12, 0.66, [{'text': name.replace('\n',''), 'size': 11.5, 'color': INK, 'bold': True, 'align': PP_ALIGN.CENTER, 'lh':1.0}])
    txt(s, l+0.08, y+1.16, bw-0.16, 0.74, [{'text': role, 'size': 9.3, 'color': MUTED, 'align': PP_ALIGN.CENTER, 'lh': 1.05}])
    if i < n-1:
        a = rect(s, l+bw-0.02, y+bh/2-0.10, 0.18, 0.20, fill=TEAL, shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
        a.rotation = 90
# specialist agents
txt(s, 0.55, 4.35, 12.2, 0.4, [{'segs':[
    ('专家 Agent（3 个） ', {'size':12,'color':AMBER,'bold':True}),
    ('—— 服务于 domain / knowledge / 交互问答等独立流程', {'size':11,'color':MUTED})]}])
spec = [
    ('Domain-Analyzer', AMBER, '业务域分析:domains→flows→steps 三层;flow-step 用 0–1 权重编码执行顺序  ·  /understand-domain'),
    ('Article-Analyzer', VIOLET, 'Karpathy 式 wiki 知识抽取:实体/主张/隐式关系(builds_on, contradicts, cites…)  ·  /understand-knowledge'),
    ('Knowledge-Graph-Guide', CYAN, '交互式导航/问答 Agent:用 jq 查询图谱,引导用户到 dashboard'),
]
yy = 4.78
for name, col, desc in spec:
    rect(s, 0.55, yy, 12.25, 0.66, fill=BG2, line=col, lw=1.1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    txt(s, 0.78, yy, 3.1, 0.66, [{'text': name, 'size': 11.5, 'color': col, 'bold': True}], anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 3.95, yy, 8.7, 0.66, [{'text': desc, 'size': 10.3, 'color': INK, 'lh':1.05}], anchor=MSO_ANCHOR.MIDDLE)
    yy += 0.74

# ============================================================================
# SLIDE 8 — Parsing engine
# ============================================================================
s = slide(); header(s, 'PARSING ENGINE', '混合解析引擎:Tree-sitter + LLM', 8)
# left: languages
lp = rect(s, 0.55, 1.62, 6.0, 5.1, fill=BG2, line=BG3, lw=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
txt(s, 0.80, 1.78, 5.5, 0.5, [{'segs':[('web-tree-sitter ', {'size':13.5,'color':TEAL,'bold':True}),('^0.26.6  (WASM 运行时)', {'size':11,'color':MUTED})]}])
txt(s, 0.80, 2.22, 5.5, 0.4, [{'text': '12 种语言的 tree-sitter 语法:', 'size': 11, 'color': MUTED}])
langs = ['JavaScript','TypeScript','Python','Rust','Go','Java','C++','C#','PHP','Ruby','Kotlin','Dart (WASM 自建包)']
cols2 = 3
for i,lg in enumerate(langs):
    r_, c_ = divmod(i, cols2)
    chip(s, 0.80 + c_*1.88, 2.62 + r_*0.52, 1.78, 0.42, lg, CYAN, fontsize=9.8)
txt(s, 0.80, 5.05, 5.5, 1.6, [
    {'text': '其它关键依赖', 'size': 12, 'color': AMBER, 'bold': True},
    {'segs':[('fuse.js ', {'size':11,'color':INK,'bold':True}),('模糊搜索  ·  ', {'size':11,'color':MUTED}),('zod ', {'size':11,'color':INK,'bold':True}),('schema 校验', {'size':11,'color':MUTED})], 'sb':7},
    {'segs':[('graphology + Louvain ', {'size':11,'color':INK,'bold':True}),('图结构/社区聚类', {'size':11,'color':MUTED})], 'sb':6},
    {'segs':[('yaml · ignore ', {'size':11,'color':INK,'bold':True}),('配置与 .understandignore', {'size':11,'color':MUTED})], 'sb':6},
])
# right: deterministic scripts + significance rules
rp = rect(s, 6.80, 1.62, 6.0, 5.1, fill=BG2, line=BG3, lw=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
txt(s, 7.05, 1.78, 5.5, 5.0, [
    {'text': '确定性脚本(事实层)', 'size': 13.5, 'color': TEAL, 'bold': True},
    {'text': 'scan-project.mjs  ·  extract-structure.mjs', 'size': 10.5, 'color': INK, 'bullet':'▸', 'sb': 8, 'lh':1.1},
    {'text': 'extract-import-map.mjs  ·  compute-batches.mjs', 'size': 10.5, 'color': INK, 'bullet':'▸', 'sb': 5, 'lh':1.1},
    {'text': 'build-fingerprints.mjs  ·  generate-ignore.mjs', 'size': 10.5, 'color': INK, 'bullet':'▸', 'sb': 5, 'lh':1.1},
    {'text': 'merge-batch-graphs.py  ·  merge-subdomain-graphs.py', 'size': 10.5, 'color': INK, 'bullet':'▸', 'sb': 5, 'lh':1.1},
    {'text': 'LLM 合成的「护栏」', 'size': 13.5, 'color': VIOLET, 'bold': True, 'sb': 16},
    {'text': '每条 import 仅生成 1 条 imports 边(来自 batchImportData)', 'size': 10.5, 'color': INK, 'bullet':'▸', 'sb': 8, 'lh':1.1},
    {'text': '显著性过滤:函数 ≥10 行 / 导出项 / 文件 >200 行', 'size': 10.5, 'color': INK, 'bullet':'▸', 'sb': 5, 'lh':1.1},
    {'text': '严格 ID 前缀:file: / function: / class: / config:', 'size': 10.5, 'color': INK, 'bullet':'▸', 'sb': 5, 'lh':1.1},
    {'text': '「信任脚本,不要重读源码」—— 降低幻觉与成本', 'size': 10.5, 'color': INK, 'bullet':'▸', 'sb': 5, 'lh':1.1},
])

# ============================================================================
# SLIDE 9 — Data model
# ============================================================================
s = slide(); header(s, 'DATA MODEL', '知识图谱数据模型:knowledge-graph.json', 9)
txt(s, 0.55, 1.52, 12.2, 0.4, [{'segs':[
    ('根结构: ', {'size':11.5,'color':TEAL,'bold':True}),
    ('version · kind("codebase"|"knowledge") · project · nodes[] · edges[] · layers[] · tour[]   (类型定义 types.ts,Zod 校验 schema.ts)', {'size':11,'color':MUTED})]}])
# node table
txt(s, 0.55, 2.00, 6.0, 0.35, [{'text': 'GraphNode 字段', 'size': 12.5, 'color': CYAN, 'bold': True}])
nrows = [['字段','类型 / 说明'],
         ['id / type / name','唯一 ID · NodeType · 名称'],
         ['filePath, lineRange','源位置 [start,end]'],
         ['summary / tags','LLM 摘要 · 主题标签[]'],
         ['complexity','simple | moderate | complex'],
         ['languageNotes','语言概念提示(可选)'],
         ['domainMeta / knowledgeMeta','业务域 / 知识库扩展(可选)']]
styled_table(s, 0.55, 2.38, 6.0, nrows, [2.25, 3.75], fontsize=9.8, row_h=0.40, head_h=0.36, header_color=CYAN)
# edge table
txt(s, 6.85, 2.00, 6.0, 0.35, [{'text': 'GraphEdge 字段', 'size': 12.5, 'color': TEAL, 'bold': True}])
erows = [['字段','类型 / 说明'],
         ['source / target','两端节点 ID'],
         ['type','EdgeType(35 种关系)'],
         ['direction','forward | backward | bidirectional'],
         ['weight','0.0 – 1.0(也用于编码顺序)'],
         ['description','关系说明(可选)']]
styled_table(s, 6.85, 2.38, 5.95, erows, [2.0, 3.95], fontsize=9.8, row_h=0.40, head_h=0.36, header_color=TEAL)
# type counts
yb = 5.35
rect(s, 0.55, yb, 12.25, 1.35, fill=BG2, line=BG3, lw=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
txt(s, 0.80, yb+0.12, 12.0, 0.4, [{'segs':[
    ('类型系统  ', {'size':12.5,'color':AMBER,'bold':True}),
    ('—— types.ts 定义 21 种节点类型 / 35 种边类型;graph-reviewer 强制校验其中核心 16 节点 / 29 边子集', {'size':10.8,'color':MUTED})]}])
txt(s, 0.80, yb+0.52, 12.0, 0.35, [{'text': '节点类型(摘选): file · function · class · module · service · endpoint · table · schema · domain · flow · step · concept · article · entity · claim · source', 'size': 10, 'color': INK, 'lh':1.1}])
txt(s, 0.80, yb+0.90, 12.0, 0.4, [{'text': '边类型(摘选): imports · calls · contains · inherits · implements · reads_from · writes_to · depends_on · tested_by · cites · contradicts · builds_on', 'size': 10, 'color': INK, 'lh':1.1}])

# ============================================================================
# SLIDE 10 — Dashboard
# ============================================================================
s = slide(); header(s, 'VISUALIZATION', '可视化 Dashboard:技术栈与渲染', 10)
# stack chips
txt(s, 0.55, 1.58, 12.2, 0.35, [{'text': '前端技术栈', 'size': 12.5, 'color': GREEN, 'bold': True}])
stack = [('React 19',GREEN),('Vite 6',GREEN),('Tailwind 4',GREEN),('Zustand 5',GREEN),('Vitest 3',GREEN),
         ('@xyflow/react (React Flow)',CYAN),('elkjs (ELK 布局)',CYAN),('graphology + Louvain',AMBER),
         ('d3-force',MUTED),('react-markdown',VIOLET),('prism-react-renderer',VIOLET)]
x = 0.55; yrow = 1.98
for label, col in stack:
    w = 0.32 + len(label)*0.092
    if x + w > 12.8: x = 0.55; yrow += 0.52
    chip(s, x, yrow, w, 0.42, label, col, fontsize=10)
    x += w + 0.14
# rendering explanation
rp = rect(s, 0.55, 3.30, 6.05, 3.40, fill=BG2, line=BG3, lw=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
txt(s, 0.80, 3.46, 5.55, 3.2, [
    {'text': '渲染机制(一个重要细节)', 'size': 12.5, 'color': CYAN, 'bold': True},
    {'text': '主图 GraphView.tsx 并非用力导向模拟布点,而是用 React Flow 渲染 + ELK 计算确定性层次布局:', 'size': 11, 'color': INK, 'sb': 8, 'lh':1.15},
    {'text': '阶段 1:容器作为原子单元,在「层」级别布局', 'size': 10.8, 'color': MUTED, 'bullet':'▸', 'sb': 8, 'lh':1.1},
    {'text': '阶段 2:展开容器时,子节点二次 ELK 布局', 'size': 10.8, 'color': MUTED, 'bullet':'▸', 'sb': 5, 'lh':1.1},
    {'text': 'Zustand 管理图状态 + 容器布局缓存(避免重算)', 'size': 10.8, 'color': MUTED, 'bullet':'▸', 'sb': 5, 'lh':1.1},
    {'text': 'd3-force / graphology 作依赖存在,主要用于聚类(具体调用点未完全确认)', 'size': 10.2, 'color': FAINT, 'bullet':'▸', 'sb': 5, 'lh':1.1, 'italic':True},
])
# features
fp = rect(s, 6.75, 3.30, 6.05, 3.40, fill=BG2, line=BG3, lw=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
txt(s, 7.00, 3.46, 5.55, 3.2, [
    {'text': '看板能力', 'size': 12.5, 'color': GREEN, 'bold': True},
    {'text': '架构层着色(API/Service/Data/UI/Utility),可平移缩放', 'size': 10.8, 'color': INK, 'bullet':'▸', 'sb': 8, 'lh':1.1},
    {'text': '模糊 + 语义搜索;点击节点看代码、关系、解释', 'size': 10.8, 'color': INK, 'bullet':'▸', 'sb': 6, 'lh':1.1},
    {'text': '依赖顺序自动生成的「导览」', 'size': 10.8, 'color': INK, 'bullet':'▸', 'sb': 6, 'lh':1.1},
    {'text': 'diff 影响叠加(ripple);业务域 / 知识图谱视图', 'size': 10.8, 'color': INK, 'bullet':'▸', 'sb': 6, 'lh':1.1},
    {'text': 'Persona 自适应 UI:junior dev / PM / power user 不同详略', 'size': 10.8, 'color': INK, 'bullet':'▸', 'sb': 6, 'lh':1.1},
    {'text': '多语言 UI:EN / 中 / 日 / 韩 / 俄 / 西 / 土', 'size': 10.8, 'color': INK, 'bullet':'▸', 'sb': 6, 'lh':1.1},
])

# ============================================================================
# SLIDE 11 — Engineering & distribution
# ============================================================================
s = slide(); header(s, 'ENGINEERING', '工程化:增量更新 · 多平台分发', 11)
boxes = [
    ('增量更新机制', TEAL, [
        'fingerprint.ts:代码身份/哈希,识别真正改动的文件',
        'staleness.ts:对比 meta.json 存储的 commit 与 HEAD',
        'change-classifier.ts:对改动分类;仅重分析变更文件',
        '为大仓库省下大量重复 LLM 调用',
    ]),
    ('自动维护 Hooks', AMBER, [
        'PostToolUse(匹配 Bash):检测 git commit/merge/rebase',
        '若启用 autoUpdate 且图谱存在 → 静默增量更新',
        'SessionStart:会话启动比对 commit,过期则刷新',
        '配置:--auto-update 开启 post-commit 钩子',
    ]),
    ('多平台分发(护城河)', CYAN, [
        '原生 /plugin install(Claude Code marketplace)',
        'install.sh / install.ps1 一行安装 + 符号链接',
        '各平台靠 .{platform}-plugin/plugin.json 自动发现',
        '覆盖 Cursor·Copilot·Codex·Gemini·Cline·Trae·Kiro… 15+',
    ]),
    ('团队协作模式', GREEN, [
        '提交整个 .understand-anything/ 目录到 Git',
        '排除 intermediate/ 与 diff-overlay.json',
        '队友复用图谱,跳过整条流水线',
        '大图谱(10MB+)用 git-lfs 跟踪',
    ]),
]
for i,(title, col, items) in enumerate(boxes):
    r_, c_ = divmod(i, 2)
    l = 0.55 + c_*6.25; t = 1.62 + r_*2.62
    rect(s, l, t, 5.95, 2.42, fill=BG2, line=col, lw=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    txt(s, l+0.25, t+0.16, 5.5, 0.4, [{'text': title, 'size': 13, 'color': col, 'bold': True}])
    lines = [{'text': it, 'size': 10.6, 'color': INK, 'bullet':'▸', 'sb': 7 if k else 2, 'lh':1.08} for k,it in enumerate(items)]
    txt(s, l+0.30, t+0.62, 5.4, 1.75, lines)

# ============================================================================
# SLIDE 12 — Commands / user workflow
# ============================================================================
s = slide(); header(s, 'COMMANDS', '命令体系与用户工作流', 12)
rows = [
    ['命令', '功能'],
    ['/understand', '分析代码库,生成知识图谱(支持增量、--language、--full 重建)'],
    ['/understand-dashboard', '启动交互式 Web 看板(Vite,token 访问)'],
    ['/understand-chat', '基于图谱对代码库提问'],
    ['/understand-explain', '深挖某个文件 / 函数 / 模块'],
    ['/understand-diff', '分析 git diff / PR:改了什么、影响哪些组件、风险'],
    ['/understand-onboard', '生成新人 onboarding 指南(可存 docs/ONBOARDING.md)'],
    ['/understand-domain', '抽取业务域,生成交互式领域流程图'],
    ['/understand-knowledge', '分析 Karpathy 式 wiki 知识库'],
]
styled_table(s, 0.55, 1.62, 7.55, rows, [2.55, 5.0], fontsize=10, row_h=0.43, head_h=0.38)
# workflow steps on right
txt(s, 8.35, 1.62, 4.45, 0.4, [{'text': '典型工作流', 'size': 12.5, 'color': TEAL, 'bold': True}])
steps = [('安装','/plugin install 或 curl install.sh'),
         ('分析','/understand → knowledge-graph.json'),
         ('探索','/understand-dashboard 可视化'),
         ('问答','/understand-chat 提问'),
         ('评审','/understand-diff 看改动影响'),
         ('维护','--auto-update 增量更新')]
yy = 2.10
for i,(t,d) in enumerate(steps):
    rect(s, 8.35, yy, 0.42, 0.42, fill=TEAL, shape=MSO_SHAPE.OVAL)
    txt(s, 8.35, yy, 0.42, 0.42, [{'text': str(i+1), 'size': 13, 'color': BG, 'bold': True, 'align': PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 8.92, yy-0.02, 3.9, 0.5, [{'segs':[(t+'  ', {'size':11.5,'color':INK,'bold':True}),(d, {'size':9.8,'color':MUTED})]}], anchor=MSO_ANCHOR.MIDDLE)
    if i < len(steps)-1:
        rect(s, 8.55, yy+0.42, 0.02, 0.34, fill=BG3)
    yy += 0.76

# ============================================================================
# SLIDE 13 — Market & competitors
# ============================================================================
s = slide(); header(s, 'MARKET LANDSCAPE', '市场定位与竞品对比', 13)
txt(s, 0.55, 1.52, 12.2, 0.4, [{'text': '「AI 代码理解 / onboarding / 代码知识图谱」赛道 2026 爆发:多款工具数月内冲上数万星。UA 处于「Agent 上下文供给」与「人类理解工具」的交叉点。', 'size': 11, 'color': MUTED, 'lh':1.1}])
rows = [
    ['竞品', '类型', '与 Understand-Anything 的差异'],
    ['CodeGraph', '本地 MCP 代码图(SQLite)', '最直接对手(~47k★),主打降低 Agent token/工具调用;偏「检索效率」'],
    ['GitNexus', '浏览器内 Graph-RAG + MCP', '纯客户端、MCP 集成深、blast-radius;但 noncommercial 许可'],
    ['DeepWiki', '托管式 AI wiki + 问答', '零配置 SaaS,私库需付费;UA 是本地、产出可拥有的图谱'],
    ['grepai', 'Agent 检索层', '主打 token 削减(声称~97%);非可视化探索产品'],
    ['Greptile', 'AI 代码评审(PR bot)', '建图但服务自动 PR 评审/查 bug,买家不同'],
    ['CodeSee / Sourcetrail', '可视化依赖图', 'CodeSee 偏图但无 LLM 问答;Sourcetrail 为 AI 前时代祖先(已归档)'],
    ['repomix / Code2Prompt', '仓库打包成单文件', '只喂上下文,无图/无 UI;与 UA 互补而非竞争'],
]
styled_table(s, 0.55, 2.02, 12.25, rows, [2.35, 3.0, 6.90], fontsize=9.8, row_h=0.55, head_h=0.40, header_color=AMBER)
txt(s, 0.55, 6.45, 12.2, 0.5, [{'segs':[
    ('最贴身竞争集: ', {'size':11.5,'color':TEAL,'bold':True}),
    ('CodeGraph · GitNexus · grepai(2026 同期梯队) + DeepWiki(最知名托管替代)。结构问题靠图谱、语义召回靠 embedding、批量上下文靠打包器——多数团队最终是混用。', {'size':11,'color':INK})]}])

# ============================================================================
# SLIDE 14 — Differentiation
# ============================================================================
s = slide(); header(s, 'DIFFERENTIATION', '差异化优势:它凭什么不一样', 14)
diffs = [
    ('① 多平台插件分发', CYAN, '15+ AI 编码平台原生覆盖,广度超过多数对手(CodeGraph 约 8 个集成)。最强护城河。'),
    ('② Tree-sitter + LLM 混合', TEAL, '确定性结构 + LLM 语义,经多 Agent 流水线产出——兼顾可信与可懂。'),
    ('③ 可提交的图谱产物', AMBER, '版本化 JSON 随仓库走,「图谱即代码」;对比托管式 SaaS 是数据所有权优势。'),
    ('④ Persona 自适应 UX', VIOLET, '按角色调详略,面向「人的理解/onboarding」——多数 Agent 检索向对手不强调。'),
    ('⑤ 人本理解框架', GREEN, '导览 + diff 影响 + 架构分层 + 语言概念提示;"教育>炫技"。'),
    ('⑥ MIT 宽松许可', PINK, '相比 GitNexus 的 noncommercial,商用更友好。'),
]
for i,(title, col, desc) in enumerate(diffs):
    r_, c_ = divmod(i, 2)
    l = 0.55 + c_*6.25; t = 1.68 + r_*1.72
    rect(s, l, t, 5.95, 1.55, fill=BG2, line=col, lw=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    rect(s, l, t, 0.12, 1.55, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    txt(s, l+0.28, t+0.16, 5.55, 0.4, [{'text': title, 'size': 13, 'color': col, 'bold': True}])
    txt(s, l+0.30, t+0.60, 5.5, 0.9, [{'text': desc, 'size': 11, 'color': INK, 'lh':1.13}])

# ============================================================================
# SLIDE 15 — Risks
# ============================================================================
s = slide(); header(s, 'RISKS & LIMITATIONS', '风险与局限:冷静看待', 15)
risks = [
    ('厂商极新、身份单薄', RED, '组织仅 1 公开仓库、~48 followers;团队/主体/融资未披露;affiliate@ 暗示 OSS 背后有商业动机。'),
    ('星标增速需警惕', AMBER, '3 个月 ~63k★,赛道内此类增速曾遭质疑;应作热度信号而非采纳证据。'),
    ('LLM 环节成本与漂移', AMBER, '语义层依赖 LLM,大仓库 token 成本高、跨次运行语义可能漂移(靠 --review / Reviewer 缓解)。'),
    ('多语言解析覆盖有限', CYAN, 'tree-sitter 需逐语言支持(如 Dart 单独建 WASM 包),小众语言能力可能不全。'),
    ('大图谱性能', CYAN, '官方提示 10MB+ 需 git-lfs;超大节点量下前端渲染性能待验证。'),
    ('图谱新鲜度依赖纪律', VIOLET, '虽有 auto-update,但团队若不规范更新,提交的图谱可能与代码脱节。'),
]
for i,(title, col, desc) in enumerate(risks):
    t = 1.66 + i*0.85
    rect(s, 0.55, t, 12.25, 0.74, fill=BG2, line=col, lw=1.1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    rect(s, 0.55, t, 0.12, 0.74, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    txt(s, 0.82, t, 3.35, 0.74, [{'text': title, 'size': 12, 'color': col, 'bold': True}], anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 4.25, t, 8.4, 0.74, [{'text': desc, 'size': 10.5, 'color': INK, 'lh':1.08}], anchor=MSO_ANCHOR.MIDDLE)

# ============================================================================
# SLIDE 16 — SWOT
# ============================================================================
s = slide(); header(s, 'SWOT', 'SWOT 洞察总结', 16)
quad = [
    ('S  优势 Strengths', GREEN, [
        '混合架构可信、产物可提交可共享',
        '15+ 平台分发广度领先',
        'MIT 许可 + 人本 UX(persona/导览)',
    ]),
    ('W  劣势 Weaknesses', RED, [
        '厂商新、身份与团队深度不明',
        'LLM 成本/漂移、大图谱性能未验证',
        '生产采纳证据不足',
    ]),
    ('O  机会 Opportunities', CYAN, [
        '赛道 2026 高速增长,双买家需求旺',
        'Agent token 降本是头号卖点',
        '可向企业版/团队协作变现',
    ]),
    ('T  威胁 Threats', AMBER, [
        'CodeGraph/GitNexus/DeepWiki 贴身竞争',
        '大厂(GitHub/Sourcegraph)可下场',
        '热度回落 / 信任问题风险',
    ]),
]
for i,(title, col, items) in enumerate(quad):
    r_, c_ = divmod(i, 2)
    l = 0.55 + c_*6.25; t = 1.66 + r_*2.55
    rect(s, l, t, 5.95, 2.35, fill=BG2, line=col, lw=1.3, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    rect(s, l, t, 5.95, 0.55, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
    txt(s, l+0.25, t+0.09, 5.5, 0.4, [{'text': title, 'size': 13, 'color': BG, 'bold': True}])
    lines = [{'text': it, 'size': 11.2, 'color': INK, 'bullet':'▸', 'sb': 9 if k else 4, 'lh':1.12} for k,it in enumerate(items)]
    txt(s, l+0.30, t+0.72, 5.4, 1.55, lines)

# ============================================================================
# SLIDE 17 — Recommendations & sources
# ============================================================================
s = slide(); header(s, 'RECOMMENDATIONS', '结语:适用建议与参考来源', 17)
# recommendations
txt(s, 0.55, 1.62, 12.2, 0.4, [{'text': '该不该用?分场景建议', 'size': 13, 'color': TEAL, 'bold': True}])
recs = [
    ('✓ 推荐试用', GREEN, '团队 onboarding、PR 影响评估、为 Claude Code/Cursor 等 Agent 提供结构化上下文——本地、MIT、可提交,低风险高收益。'),
    ('◐ 谨慎评估', AMBER, '作为企业关键路径的唯一依赖前,先验证大仓库性能、LLM 成本与图谱维护流程;关注厂商可持续性。'),
    ('✗ 暂不适合', RED, '需要强 SLA/合规背书、或主体明确的供应商关系时;小众语言深度解析需求。'),
]
for i,(tag, col, desc) in enumerate(recs):
    t = 2.05 + i*0.84
    rect(s, 0.55, t, 12.25, 0.74, fill=BG2, line=col, lw=1.1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    txt(s, 0.78, t, 2.4, 0.74, [{'text': tag, 'size': 12.5, 'color': col, 'bold': True}], anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 3.15, t, 9.5, 0.74, [{'text': desc, 'size': 10.8, 'color': INK, 'lh':1.08}], anchor=MSO_ANCHOR.MIDDLE)
# sources
rect(s, 0.55, 4.72, 12.25, 1.95, fill=BG2, line=BG3, lw=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
txt(s, 0.80, 4.86, 11.8, 1.75, [
    {'text': '参考来源', 'size': 12, 'color': CYAN, 'bold': True},
    {'text': 'github.com/Egonex-AI/Understand-Anything (源码 + GitHub REST API,权威:stars/forks/license/日期)  ·  understand-anything.com', 'size': 9.8, 'color': MUTED, 'sb': 6, 'lh':1.2},
    {'text': '竞品:CodeGraph (colbymchenry) · GitNexus (abhigyanpatwari) · DeepWiki (cognition.ai) · Greptile · Sourcegraph Cody · CodeSee · GitDiagram · repomix', 'size': 9.8, 'color': MUTED, 'sb': 5, 'lh':1.2},
    {'segs':[('数据可信度: ', {'size':9.5,'color':AMBER,'bold':True}),
             ('stars/forks/license/日期经 API 校验;版本数、竞品星数、市场规模为二手来源(方向性);厂商团队/融资未披露。', {'size':9.5,'color':FAINT})], 'sb': 7, 'lh':1.2},
    {'text': '本报告由 Claude Code 以 workflow 多 Agent 调研模式产出 · 2026-06-18', 'size': 9, 'color': FAINT, 'sb': 8},
])

# ----------------------------------------------------------------------------
out = '/home/user/openai-chatgpt-codex/reports/Understand-Anything-深度洞察报告.pptx'
prs.save(out)
print('SAVED:', out)
print('Slides:', len(prs.slides._sldIdLst))
