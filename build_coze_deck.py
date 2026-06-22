# -*- coding: utf-8 -*-
"""Build the ByteDance Coze (扣子) deep-insight deck (PPTX)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
import os

# ---------- palette ----------
NAVY   = RGBColor(0x0C, 0x16, 0x2C)   # deep background
INK    = RGBColor(0x1F, 0x29, 0x37)   # body text
BLUE   = RGBColor(0x2B, 0x6C, 0xF6)   # 扣子 / accent
CYAN   = RGBColor(0x12, 0xB5, 0xCB)
PURPLE = RGBColor(0x6E, 0x4A, 0xF0)
GREEN  = RGBColor(0x12, 0x9A, 0x5E)
GRAY   = RGBColor(0x6B, 0x72, 0x80)
LGRAY  = RGBColor(0xE5, 0xE7, 0xEB)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
AMBER  = RGBColor(0xC2, 0x6B, 0x09)
RED    = RGBColor(0xB9, 0x1C, 0x1C)
CARD   = RGBColor(0xF3, 0xF6, 0xFB)

FONT = "Microsoft YaHei"
CHARTS = "/home/user/openai-chatgpt-codex/coze-insight/charts"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

def set_font(run, name=FONT):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ('a:latin', 'a:ea', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set('typeface', name)

def rect(slide, x, y, w, h, color, line=None):
    sp = slide.shapes.add_shape(1, x, y, w, h)  # rectangle
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp

def textbox(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    return tb, tf

def style(p, text, size, color=INK, bold=False, align=PP_ALIGN.LEFT, space_after=4, italic=False):
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(0)
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color
    set_font(r)
    return r

def header(slide, kicker, title):
    rect(slide, 0, 0, SW, Inches(1.18), NAVY)
    rect(slide, 0, Inches(1.18), SW, Pt(3), BLUE)
    tb, tf = textbox(slide, Inches(0.55), Inches(0.16), Inches(12.2), Inches(0.95))
    p = tf.paragraphs[0]; style(p, kicker, 11, BLUE, bold=True, space_after=2)
    p2 = tf.add_paragraph(); style(p2, title, 24, WHITE, bold=True)

def footer(slide, n, note=""):
    tb, tf = textbox(slide, Inches(0.55), Inches(7.06), Inches(11.0), Inches(0.36))
    p = tf.paragraphs[0]
    style(p, note, 8, GRAY)
    tb2, tf2 = textbox(slide, Inches(12.4), Inches(7.06), Inches(0.7), Inches(0.36))
    p2 = tf2.paragraphs[0]
    style(p2, str(n), 9, GRAY, align=PP_ALIGN.RIGHT)

def content_slide(kicker, title, n, note="可信度: 强=官方+多源 · 中=单一媒体/厂商自报 · 弱=已剔除/降级 | 聚焦国内版 coze.cn · 数据截至 2026-06"):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, SW, SH, WHITE)
    header(s, kicker, title)
    footer(s, n, note)
    return s

def bullets(slide, items, x=Inches(0.6), y=Inches(1.5), w=Inches(12.1), h=Inches(5.4),
            size=14, gap=7):
    tb, tf = textbox(slide, x, y, w, h)
    first = True
    for it in items:
        lvl = it.get("lvl", 0)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = lvl
        p.space_after = Pt(it.get("gap", gap)); p.space_before = Pt(0)
        p.alignment = PP_ALIGN.LEFT
        marks = {0: "▍ ", 1: "•  ", 2: "–  "}
        prefix = it.get("bullet", marks.get(lvl, "•  "))
        col = it.get("color", INK if lvl == 0 else GRAY if lvl >= 2 else INK)
        sz = it.get("size", size if lvl == 0 else size-2)
        rp = p.add_run(); rp.text = prefix
        rp.font.size = Pt(sz); rp.font.bold = (lvl == 0)
        rp.font.color.rgb = (BLUE if lvl == 0 else col)
        set_font(rp)
        if "lead" in it:
            rl = p.add_run(); rl.text = it["lead"]
            rl.font.size = Pt(sz); rl.font.bold = True; rl.font.color.rgb = col
            set_font(rl)
        rt = p.add_run(); rt.text = it["t"]
        rt.font.size = Pt(sz); rt.font.bold = it.get("bold", lvl == 0 and "lead" not in it)
        rt.font.color.rgb = col
        set_font(rt)
        if "tag" in it:
            rg = p.add_run(); rg.text = "  〔" + it["tag"] + "〕"
            rg.font.size = Pt(sz-2); rg.font.italic = True
            rg.font.color.rgb = it.get("tagc", GRAY)
            set_font(rg)
    return tb

def table_slide(kicker, title, n, headers, rows, col_w, note="", fs=11, hfs=11,
                y=Inches(1.45), row_h=Inches(0.42)):
    s = content_slide(kicker, title, n, note)
    nrows, ncols = len(rows)+1, len(headers)
    total_w = sum(col_w)
    x = Inches((13.333 - total_w)/2)
    gtab = s.shapes.add_table(nrows, ncols, x, y, Inches(total_w), row_h*nrows)
    tbl = gtab.table
    tbl.first_row = True; tbl.horz_banding = True
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for j, htext in enumerate(headers):
        c = tbl.cell(0, j)
        c.fill.solid(); c.fill.fore_color.rgb = NAVY
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        c.margin_left = Inches(0.07); c.margin_right = Inches(0.05)
        c.margin_top = Inches(0.02); c.margin_bottom = Inches(0.02)
        p = c.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
        r = p.add_run(); r.text = htext; r.font.size = Pt(hfs); r.font.bold = True
        r.font.color.rgb = WHITE; set_font(r)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.fill.solid(); c.fill.fore_color.rgb = WHITE if i % 2 else CARD
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.margin_left = Inches(0.07); c.margin_right = Inches(0.05)
            c.margin_top = Inches(0.02); c.margin_bottom = Inches(0.02)
            p = c.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
            txt = val; col = INK; bold = (j == 0)
            if isinstance(val, tuple):
                txt, col = val[0], val[1]
                bold = val[2] if len(val) > 2 else (j == 0)
            r = p.add_run(); r.text = txt; r.font.size = Pt(fs)
            r.font.bold = bold; r.font.color.rgb = col; set_font(r)
    return s

def img(slide, name, x, y, w=None, h=None):
    path = os.path.join(CHARTS, name)
    kw = {}
    if w is not None: kw["width"] = w
    if h is not None: kw["height"] = h
    return slide.shapes.add_picture(path, x, y, **kw)

def chip(slide, x, y, w, h, title, body, accent):
    rect(slide, x, y, w, h, CARD)
    rect(slide, x, y, Inches(0.09), h, accent)
    tb, tf = textbox(slide, x+Inches(0.22), y+Inches(0.12), w-Inches(0.36), h-Inches(0.2))
    p = tf.paragraphs[0]; style(p, title, 13, accent, bold=True, space_after=4)
    for ln in body:
        pp = tf.add_paragraph(); style(pp, ln, 11, INK, space_after=2)

# ============================================================ SLIDE 1 — COVER
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, SW, SH, NAVY)
rect(s, 0, Inches(4.05), SW, Pt(3.5), BLUE)
rect(s, Inches(0.0), 0, Inches(0.22), SH, BLUE)
tb, tf = textbox(s, Inches(0.9), Inches(1.4), Inches(11.7), Inches(2.7))
p = tf.paragraphs[0]; style(p, "字节跳动 · 扣子 COZE", 17, BLUE, bold=True, space_after=6)
p = tf.add_paragraph(); style(p, "深度洞察分析", 46, WHITE, bold=True, space_after=4)
p = tf.add_paragraph(); style(p, "Product · Technology · Market · Business · Ecosystem · Competition", 15, LGRAY)
tb, tf = textbox(s, Inches(0.9), Inches(4.35), Inches(11.7), Inches(2.3))
p = tf.add_paragraph(); style(p, "综合全景：产品定位 · 功能架构 · 底层技术与开源 · 用户与市场 · 商业化变现 · 生态协同 · 竞争格局 · 风险机会", 13, LGRAY, space_after=10)
p = tf.add_paragraph(); style(p, "聚焦国内版（coze.cn / 扣子）   |   2026 年 6 月   |   5 路并行检索 + 多源交叉验证 + 可信度分层标注", 12, BLUE, bold=True, space_after=4)
p = tf.add_paragraph(); style(p, "方法学提示：本次环境对多数中文媒体/官网 WebFetch 受限(403)，多数结论基于权威页面搜索摘要并经多源交叉印证；GitHub 数据为 API 实测一手值。引用前建议对照官网/官方发布稿核对确切数字与日期。", 9.5, GRAY)

# ============================================================ SLIDE 2 — EXEC SUMMARY
s = content_slide("EXECUTIVE SUMMARY", "执行摘要：六个核心洞察", 2)
bullets(s, [
 {"lead":"定位演进，而非一成不变 — ", "t":"扣子是字节基于自研豆包大模型 + 火山引擎底座的一站式 AI Agent 开发平台；两年内从“零代码 Bot 工厂”→“通用 Agent(扣子空间)”→“职场 AI + Vibe Coding(2.0)”，定位多次升级。", "tag":"强"},
 {"lead":"规模领先，但要看口径 — ", "t":"月活开发者 100万(2024.12)→ 300万(2025.12)；2.0 时官方称服务“上千万用户 / 1000万真实开发场景”，自称“中国最大的 Agent 开发平台”。口径随时间演变，勿混为单一曲线。", "tag":"中-强"},
 {"lead":"开源是关键战略落子 — ", "t":"2025.7.26 开源 Coze Studio + Coze Loop(Apache 2.0，可商用)；Studio 已达约 2.1 万 Star。意在让扣子成为 Agent 开发“事实标准”、正面反制 Dify。", "tag":"强"},
 {"lead":"商业化“三段式火箭” — ", "t":"四档订阅 + 资源点按量计费 + 模板分成(平台抽 20%)；“扣子(开发者/C 端) + HiAgent(企业私有化)”双平台；开源引流 → 企业订阅 → 火山引擎算力/抖音流量变现。", "tag":"中-强"},
 {"lead":"护城河=模型×算力×流量×开源 — ", "t":"豆包(日均 token 两年增约 1000 倍至 120 万亿)、火山引擎(模型调用份额升至 59.2%)、抖音/飞书分发生态、开源社区四位一体。", "tag":"中-强"},
 {"t":"关键风险：工具型变现天花板与低粘性、定位反复摇摆、与腾讯/百度/阿里/Dify 高度同质化、母公司利润波动加大算力投入压力。", "lvl":1, "color":AMBER, "bullet":"⚠  "},
], y=Inches(1.42), size=13.5, gap=10)

# ============================================================ SLIDE 3 — METHOD & RECONCILIATION
s = content_slide("METHODOLOGY", "研究方法与关键澄清", 3)
bullets(s, [
 {"lead":"流程 — ", "t":"将问题拆为 5 个角度(产品 / 市场 / 商业 / 技术 / 竞争)并行检索 → 去重抓取 → 多源交叉核验 → 按可信度排序综合。"},
 {"lead":"可信度分层 — ", "t":"强(官方 + 多源交叉)；中(单一媒体或厂商自报、无第三方审计)；弱(已剔除或降级)。全程标注。"},
 {"t":"关键澄清①：豆包 ≠ 扣子(最易踩坑)", "lvl":1, "color":RED, "bullet":"✗  ", "bold":True},
 {"t":"豆包是 C 端 AI 助手 App(月活以亿计)，扣子是面向开发者的 Agent 开发平台(规模以开发者/智能体数计)。“5700万日活”“3.4亿月活”均为豆包，不可挂在扣子名下。", "lvl":2, "color":GRAY},
 {"t":"关键澄清②：开发者数量口径随时间变化", "lvl":1, "color":AMBER, "bullet":"~  ", "bold":True},
 {"t":"100万“活跃开发者”(2024.12)→ 300万“月活开发者”(2025.12)→ “上千万用户/1000万开发场景”(2026.01)。三者计量口径不同，构成增长趋势但非同一指标。", "lvl":2, "color":GRAY},
 {"lead":"核心限制 — ", "t":"本次 WebFetch 对 coze.cn / 火山引擎 / 主流中文媒体多数 403，结论基于搜索摘要交叉印证；GitHub(开源 Star/版本)为 API 直读一手数据。定价/单源数字建议出稿前以官网复核。", "color":AMBER, "tag":"重要"},
], y=Inches(1.4), size=12.8, gap=7)

# ============================================================ SLIDE 4 — PRODUCT POSITIONING
s = content_slide("PRODUCT 1/3 · POSITIONING", "产品定位：扣子是什么", 4)
bullets(s, [
 {"lead":"一句话 — ", "t":"无论是否有编程基础，都能快速搭建基于大模型的智能体 / AI 应用，并一键发布到多渠道的“一站式 AI Agent 开发平台”(零代码 / 低代码)。", "tag":"强"},
 {"lead":"双版本 — ", "t":"海外版 Coze(coze.com)2023.11 上线，国内版扣子(coze.cn)2024.02 上线；国内版早期内置“云雀”，现以自研豆包为主力并开放多模型。", "tag":"强"},
 {"lead":"目标用户三层 — ", "t":"① 个人创作者(私域营销 / 自媒体)；② 中小开发者；③ 企业级客户(扣子专业版 / 私有化)。订阅分免费版 / 进阶版 / 团队版 / 企业版。", "tag":"中-强"},
 {"lead":"战略洞察：“严肃开发者” — ", "t":"据极客公园深度报道，扣子团队最初想做“全民可用”(像剪映)，后聚焦到有商业目的、靠它赚钱或解决效率问题的“严肃开发者”，定位经历多次校准。", "tag":"中"},
 {"lead":"在字节版图中的角色 — ", "t":"模型(豆包)→ 应用构建(扣子)→ 分发(抖音/飞书/豆包)→ 算力(火山引擎)闭环中的“应用构建层”。", "tag":"中-强"},
], y=Inches(1.5), size=13.5, gap=12)

# ============================================================ SLIDE 5 — TIMELINE (chart)
s = content_slide("PRODUCT 2/3 · TIMELINE", "产品演进时间轴：从 Bot 工厂到职场 AI", 5,
                  note="日期与口径多源交叉 · 2025.12 与 2026.01 为相邻的两次发布(Force 大会 → 扣子 2.0 品牌升级) | 数据截至 2026-06")
img(s, "timeline.png", Inches(0.4), Inches(1.55), w=Inches(12.55))
tb, tf = textbox(s, Inches(0.6), Inches(5.7), Inches(12.1), Inches(1.2))
p = tf.paragraphs[0]
style(p, "节奏特征：上线即高速迭代——2024 完成商业化收费，2025 连推“通用 Agent(扣子空间)”与“开源(Studio/Loop)”两大动作，", 12, INK, space_after=3)
p = tf.add_paragraph()
style(p, "2025 年底至 2026 年重心转向“职场生产力 + Vibe Coding”，并为 Agent 配置邮箱 / 云电脑 / 云手机(Agent World)。", 12, INK)

# ============================================================ SLIDE 6 — FUNCTIONAL ARCHITECTURE
s = content_slide("PRODUCT 3/3 · CAPABILITIES", "核心功能架构：六大能力模块 + 多渠道发布", 6)
chip(s, Inches(0.55), Inches(1.5), Inches(3.95), Inches(1.5), "① 工作流 Workflow",
     ["可视化拖拽编排，节点含 LLM、", "知识库、插件、代码、条件分支、", "多模态、智能体节点等"], BLUE)
chip(s, Inches(4.69), Inches(1.5), Inches(3.95), Inches(1.5), "② 插件 Plugin",
     ["200+ 官方插件 + 自定义插件，", "覆盖办公/设计/数据/舆情等；", "可“MCP 化”发布"], CYAN)
chip(s, Inches(8.83), Inches(1.5), Inches(3.95), Inches(1.5), "③ 知识库 RAG",
     ["TXT/PDF/网页/Notion/飞书等", "自动向量化 + 检索，", "降低大模型幻觉"], PURPLE)
chip(s, Inches(0.55), Inches(3.15), Inches(3.95), Inches(1.5), "④ 记忆体系",
     ["变量 / 数据库 / 长期记忆，", "形成对用户的个性化记忆，", "支持读写权限控制"], GREEN)
chip(s, Inches(4.69), Inches(3.15), Inches(3.95), Inches(1.5), "⑤ 多 Agent 编排",
     ["多智能体协同，按角色分工", "(数据/分析/创作/执行)；", "触发器支持定时/事件触发"], BLUE)
chip(s, Inches(8.83), Inches(3.15), Inches(3.95), Inches(1.5), "⑥ 多模态 / 图像视频流",
     ["图像流可视化编排，", "集成 Seedance 视频生成、", "语音合成，构建“能看能听”Agent"], CYAN)
tb, tf = textbox(s, Inches(0.55), Inches(4.95), Inches(12.2), Inches(1.5))
p = tf.paragraphs[0]
style(p, "发布渠道（一键多端分发，扣子核心优势之一）", 13, NAVY, bold=True, space_after=5)
p = tf.add_paragraph()
style(p, "豆包 · 飞书 · 抖音(小程序/私信助手) · 微信(公众号/客服/小程序) · 企业微信 · API 服务 · Chat SDK · 网页/嵌入网站 · 扣子商店/模板", 12.5, INK)

# ============================================================ SLIDE 7 — TECH 1: MODELS
s = content_slide("TECH 1/2 · MODELS", "技术深钻①：底层模型与多模型开放", 7)
bullets(s, [
 {"lead":"主力模型豆包 1.6(2025.06)— ", "t":"含 seed-1.6 / thinking / flash 三款；支持深度思考开/关/自适应三模式、多模态输入、256K 上下文；首创“按输入长度区间定价”，综合成本约为豆包 1.5 或 DeepSeek-R1 的 1/3。", "tag":"强"},
 {"lead":"多模型开放(非单一锁定)— ", "t":"2024.06 上线“模型广场/竞技场”(匿名 PK 投票)，接入豆包、通义千问、智谱 GLM、MiniMax、Kimi 等；2025 初起接入 DeepSeek R1/V3，支持其思维链与 Function Calling。", "tag":"中-强"},
 {"lead":"开源版可接模型协议 — ", "t":"Coze Studio 模型配置支持火山方舟(Ark)、OpenAI、DeepSeek、Claude、Gemini、Qwen、Ollama，配置 api_key + model 即可接入。", "tag":"强(仓库直读)"},
 {"lead":"底座关系 — ", "t":"豆包经火山引擎对外提供，火山引擎提供算力与基础设施；扣子专业版即由火山引擎承载。模型→平台→算力一体。", "tag":"强"},
 {"t":"存疑：SaaS 版当前在线模型完整清单与最新版本(如 GLM-4.6 / Qwen3 / Kimi K2 是否在列)未能从官方页核实；豆包 1.6 之后是否有更新版本本轮未覆盖。", "lvl":1, "color":GRAY, "bullet":"○  "},
], y=Inches(1.5), size=13, gap=10)

# ============================================================ SLIDE 8 — TECH 2: ARCHITECTURE
s = content_slide("TECH 2/2 · ARCHITECTURE", "技术深钻②：运行时架构、RAG 与 MCP", 8)
bullets(s, [
 {"lead":"技术栈(开源版直读)— ", "t":"后端 Golang(HTTP 框架字节自研 Hertz)、前端 React + TypeScript、微服务 + 领域驱动设计(DDD)；最低 2 核 CPU + 4GB 内存即可本地部署。", "tag":"强"},
 {"lead":"核心运行时 = CloudWeGo Eino — ", "t":"字节开源的 Go 大模型应用框架，已被抖音、豆包、Coze 多业务线采用；为扣子提供 Agent/工作流运行时、模型抽象、知识库索引与检索。", "tag":"强"},
 {"lead":"RAG 检索 — ", "t":"开源版默认精确检索用 Elasticsearch、语义检索用 Milvus(可切换火山引擎云原生向量库 VikingDB)；中间件含 MySQL 8.4 / Redis 8 / etcd / MinIO / NSQ。", "tag":"强(配置直读)"},
 {"lead":"MCP 全面支持 — ", "t":"2025.04 起工作流可作为 MCP 扩展，插件/工作流均可“MCP 化”，开发者可发布自研 MCP；内置高德、飞书多维表格、GitHub、MySQL、ClickHouse、语音合成等扩展。", "tag":"中-强"},
 {"lead":"AgentOps = Coze Loop — ", "t":"覆盖 Prompt 开发(Playground/版本管理)、评测(多维自动化测试)、全链路可观测(input→output 每步记录)；经 Eino 接入多模型，提供多语言 SDK。", "tag":"强"},
], y=Inches(1.5), size=12.8, gap=9)

# ============================================================ SLIDE 9 — OPEN SOURCE (chart)
s = content_slide("OPEN-SOURCE STRATEGY", "开源战略：Coze Studio + Coze Loop", 9,
                  note="GitHub Stars/Forks 为 API 实测(2026-06-22) · 开源时间与协议多源一致 | 数据截至 2026-06")
img(s, "github_stars.png", Inches(0.5), Inches(1.5), w=Inches(6.5))
bullets(s, [
 {"lead":"时间与协议 — ", "t":"2025.7.26 开源，均为 Apache 2.0，明确允许商用。", "size":12.5},
 {"lead":"Coze Studio — ", "t":"开发平台开源版，脱胎于服务“数万企业、数百万开发者”的商业版；含工作流引擎、插件框架、知识库。", "size":12.5},
 {"lead":"Coze Loop — ", "t":"AgentOps 全生命周期管理(Prompt/评测/可观测)。", "size":12.5},
 {"lead":"开源 vs 商业 — ", "t":"核心搭建能力开源；音色/语音、租户管理、企业弹性扩容等留作商业版变现点。", "size":12.5, "color":AMBER},
 {"lead":"战略意图 — ", "t":"以开源抢占 Agent 开发“事实标准”，正面对标/反制 Dify，并为企业版与火山引擎云服务引流。", "size":12.5, "color":BLUE},
], x=Inches(7.25), y=Inches(1.6), w=Inches(5.6), size=12.5, gap=9)

# ============================================================ SLIDE 10 — MARKET GROWTH (chart)
s = content_slide("MARKET 1/2 · GROWTH", "市场地位与增长：中国最大的 Agent 开发平台", 10,
                  note="官方披露口径(Force 大会 / 扣子 2.0 发布稿) · 口径随时间演变，详见图注 | 数据截至 2026-06")
img(s, "dev_growth.png", Inches(0.45), Inches(1.5), w=Inches(7.0))
bullets(s, [
 {"lead":"自身增长 — ", "t":"月活开发者 100万(2024.12)→ 300万(2025.12)；智能体数 200万+(2024.12 起)。", "size":12.5},
 {"lead":"2.0 口径跃升 — ", "t":"2026.01 称服务“上千万用户 / 1000万真实开发场景”，自称“中国最大的 Agent 开发平台”。", "size":12.5, "color":BLUE},
 {"lead":"第三方定位 — ", "t":"多份 2025 选型榜单将扣子列为国内 Agent 平台第一梯队/榜首(用户规模最高、品牌热度最大)；量子位智库《年度 AI 100》列入 AI Agent 赛道 TOP3。", "size":12.5},
 {"t":"存疑：扣子“月活/月访问用户”有 200万 与 458万(2025.06) 两个口径(或分指 App 与网页端)；“首日 50 万”等单源数字建议复核。", "lvl":1, "color":GRAY, "bullet":"○  ", "size":11.5},
], x=Inches(7.65), y=Inches(1.55), w=Inches(5.3), size=12.5, gap=9)

# ============================================================ SLIDE 11 — MARKET SHARE & FOUNDATION (charts)
s = content_slide("MARKET 2/2 · SHARE & FOUNDATION", "市场份额与底座规模：豆包 + 火山引擎", 11,
                  note="IDC 数据为 MaaS/大模型平台口径，非 Agent 平台细分口径 · 豆包 token 为发布会披露的发展数据 | 数据截至 2026-06")
img(s, "maas_share.png", Inches(0.35), Inches(1.5), h=Inches(3.4))
img(s, "doubao_tokens.png", Inches(7.0), Inches(1.55), w=Inches(5.95))
tb, tf = textbox(s, Inches(0.55), Inches(5.2), Inches(12.3), Inches(1.7))
p = tf.paragraphs[0]
style(p, "底座护城河：豆包大模型日均 token 两年增长约 1000 倍(2024.05 的 1200 亿 → 2026.03 的 120 万亿，居中国第一、全球前三)；", 12.5, INK, space_after=4)
p = tf.add_paragraph()
style(p, "火山引擎大模型云上调用量份额从 2024 年 46.4% 升至 2025 年 59.2%(IDC)。注意：IDC MaaS 口径下字节约 16%、列第三；", 12.5, INK, space_after=4)
p = tf.add_paragraph()
style(p, "截至本研究，未发现针对“Agent 开发平台”细分赛道给出扣子的明确市占率百分比，故不应为扣子标注精确市占。", 12, AMBER)

# ============================================================ SLIDE 12 — BUSINESS MODEL (table)
table_slide(
 "BUSINESS 1/2 · PRICING", "商业模式与变现：四档订阅 + 资源点计费", 12,
 ["版本", "价格(人民币)", "资源/权益", "可信度"],
 [
  ["个人免费版", "0", "约 500 资源点/天(每日重置)", ("高", GREEN, True)],
  ["个人进阶版", "首月 9.9 元(原价 39.9 元/月)", "约 3 万积分/月", ("中", AMBER, True)],
  ["团队版", "按席位包年包月(单价未确证)", "协同开发、跨空间迁移、方舟模型接入", ("中", AMBER, True)],
  ["企业版", "约 4980 元/月，年付 8.3 折", "约 300 万资源点/月", ("中", AMBER, True)],
  ["专业版(付费)", "按调用 Token 计费", "付费用户每日赠 500 资源点", ("高", GREEN, True)],
 ],
 col_w=[2.1, 4.0, 4.3, 1.4],
 note="定价随计费规则在 2024–2026 间反复调整，时效性强；带“中”者为第三方整理，未经官网直接确认，以官网为准 | 数据截至 2026-06",
 fs=11.5, hfs=12, row_h=Inches(0.62), y=Inches(1.5))
s = prs.slides[-1]
bullets(s, [
 {"lead":"计费单位 — ", "t":"资源点(积分)为核心结算；模型费 = Token 用量 × 模型单价；扣减顺序：资源点 > 代金券 > 现金账户。智能体交互已改为仅按模型 Token 计费(不再单独按次收费)。", "size":11.5},
 {"lead":"创作者变现 — ", "t":"2024.10 上线模板中心，创作者可定价售卖智能体模板，平台抽成 20%(创作者得 80%)；以 50 万奖池赛事拉新。", "size":11.5},
], x=Inches(0.6), y=Inches(5.55), w=Inches(12.1), h=Inches(1.3), size=11.5, gap=7)

# ============================================================ SLIDE 13 — COMMERCIALIZATION STRATEGY
s = content_slide("BUSINESS 2/2 · STRATEGY", "商业化战略：双平台 + 三段式火箭", 13)
chip(s, Inches(0.55), Inches(1.55), Inches(6.0), Inches(2.0), "双平台分工",
     ["• 扣子(coze.cn)：面向 C 端/开发者，",
      "  低代码、易用、广度，含开源版",
      "• HiAgent(火山引擎)：面向企业级，",
      "  支持私有化部署、数据不出内网，",
      "  满足金融/医疗/制造合规"], BLUE)
chip(s, Inches(6.75), Inches(1.55), Inches(6.0), Inches(2.0), "企业落地“1+N+X”",
     ["• 1 个统一控制台",
      "• N 个预置通用 Agent",
      "• X 个企业自建 Agent",
      "宣称 AI 应用开发周期缩短 95%+",
      "(火山引擎厂商口径)"], PURPLE)
tb, tf = textbox(s, Inches(0.55), Inches(3.85), Inches(12.2), Inches(0.5))
style(tf.paragraphs[0], "三段式火箭（媒体/分析框架，非官方表述）", 14, NAVY, bold=True)
chip(s, Inches(0.55), Inches(4.45), Inches(3.95), Inches(1.7), "① 开源引流",
     ["社区版 Apache 2.0", "免费可商用、本地部署，", "吸引开发者、抢事实标准"], GREEN)
chip(s, Inches(4.69), Inches(4.45), Inches(3.95), Inches(1.7), "② 企业订阅变现",
     ["专属集群、SLA 保障、", "租户管理、安全合规等", "高级功能付费"], BLUE)
chip(s, Inches(8.83), Inches(4.45), Inches(3.95), Inches(1.7), "③ 生态层盈利",
     ["火山引擎算力/方舟模型", "Token 费 + 抖音流量分发", "(变现与云资源深度捆绑)"], PURPLE)
tb, tf = textbox(s, Inches(0.55), Inches(6.25), Inches(12.2), Inches(0.6))
style(tf.paragraphs[0], "线索：据极客公园报道，扣子团队 2025 年初设定的商业化目标“半年内即达成”；但无公开营收金额，属定性表述。", 11.5, GRAY, italic=True)

# ============================================================ SLIDE 14 — ECOSYSTEM MATRIX
s = content_slide("ECOSYSTEM", "字节 AI 生态协同：模型 → 应用 → 分发 → 算力闭环", 14)
chip(s, Inches(0.55), Inches(1.55), Inches(3.95), Inches(1.75), "豆包 App",
     ["C 端 AI 助手入口", "月活以亿计(2026Q1 约 3.4 亿)", "为生态导流"], BLUE)
chip(s, Inches(4.69), Inches(1.55), Inches(3.95), Inches(1.75), "即梦 / Seedance",
     ["图像 / 视频生成", "多模态创作能力", "供扣子工作流调用"], CYAN)
chip(s, Inches(8.83), Inches(1.55), Inches(3.95), Inches(1.75), "扣子 Coze",
     ["Agent / AI 应用构建层", "开发者与企业的搭建平台", "(本报告主角)"], PURPLE)
chip(s, Inches(0.55), Inches(3.45), Inches(3.95), Inches(1.75), "火山引擎 + HiAgent",
     ["算力 / 方舟模型 / 向量库", "企业级私有化 Agent", "商业化与变现底座"], GREEN)
chip(s, Inches(4.69), Inches(3.45), Inches(3.95), Inches(1.75), "飞书智能伙伴 / Aily",
     ["办公协同场景 Agent", "企业内部知识与流程", "B 端落地入口"], BLUE)
chip(s, Inches(8.83), Inches(3.45), Inches(3.95), Inches(1.75), "抖音 / 头条 / 番茄",
     ["海量 C 端流量分发", "智能体可直接部署触达", "扣子独特冷启动优势"], AMBER)
tb, tf = textbox(s, Inches(0.55), Inches(5.45), Inches(12.2), Inches(1.2))
p = tf.paragraphs[0]
style(p, "协同逻辑：豆包(模型)训练能力 → 扣子(平台)让能力可被搭建 → 抖音/飞书(渠道)分发触达 → 火山引擎(算力)承载并变现。", 12.5, INK, space_after=4)
p = tf.add_paragraph()
style(p, "这条“模型—应用—分发—算力”闭环是扣子相对独立 Agent 平台的最大结构性差异。", 12.5, BLUE, bold=True)

# ============================================================ SLIDE 15 — COMPETITION (table)
table_slide(
 "COMPETITIVE LANDSCAPE", "竞争格局对比：国内主要 Agent 平台", 15,
 ["平台", "厂商/模型", "定位与关键事实", "分发壁垒"],
 [
  ["扣子 Coze", "字节/豆包", ("零/低代码、易用、广度强；开源 Studio/Loop；个人+企业双线", INK, False), ("抖音/飞书/豆包", BLUE, True)],
  ["腾讯元器", "腾讯/混元", ("免费、偏 C 端创作；一键分发，打通微信支付/企微 CRM", INK, False), ("微信/QQ/公众号", INK, False)],
  ["腾讯云 ADP", "腾讯云/混元", ("ADP 3.0(2025.09)，偏 To B 严肃生产，强 RAG/合规", INK, False), ("企业渠道", INK, False)],
  ["文心 AgentBuilder", "百度/文心", ("零低代码；核心=百度搜索/文心 App 流量；80万开发者(2024.11)", INK, False), ("百度搜索/地图", INK, False)],
  ["阿里云百炼", "阿里/通义", ("模型服务+Agent 双核；20万+开发者、80万+ Agent；企业级", INK, False), ("阿里云/钉钉", INK, False)],
  ["Dify", "开源", ("一站式工业化流水线，企业级运维完整；扣子开源后最直接对手", INK, False), ("中立/自托管", INK, False)],
 ],
 col_w=[2.2, 1.7, 6.7, 1.7],
 note="各家定位/数字多为厂商披露或媒体整理(中-高可信)；扣子优势在“易用+流量生态”，劣势在与竞品高度同质化 | 数据截至 2026-06",
 fs=10.5, hfs=11, row_h=Inches(0.66), y=Inches(1.5))

# ============================================================ SLIDE 16 — MOATS
s = content_slide("MOATS", "护城河分析：四位一体的结构性优势", 16)
bullets(s, [
 {"lead":"① 自研模型底座 — ", "t":"豆包系列 + 区间定价把成本压到 DeepSeek-R1 约 1/3；日均 token 两年增约 1000 倍至 120 万亿，模型能力与成本双重护城。", "tag":"中-强"},
 {"lead":"② 火山引擎算力 — ", "t":"大模型云上调用量份额 2025 年 59.2%(IDC)，超百万家企业/个人使用；为扣子提供算力、推理、向量库底座并直接变现。", "tag":"中-强"},
 {"lead":"③ 字节流量分发(最独特)— ", "t":"智能体可直接部署到抖音、今日头条、飞书等字节系产品，被评“生态繁荣度行业独一无二”；解决 Agent 冷启动与触达难题。", "tag":"中"},
 {"lead":"④ 开源生态 — ", "t":"Apache 2.0 开源 + 2核4G 本地部署，抢占开发者心智与“事实标准”，对标 Dify；社区贡献反哺产品。", "tag":"强"},
 {"lead":"⑤ 易用性 + 多模型开放 — ", "t":"可视化拖拽 + 200+ 插件 + 多模型(豆包/DeepSeek/Qwen/GLM…)，降低门槛、避免单一模型锁定担忧。", "tag":"中"},
], y=Inches(1.5), size=13.5, gap=11)

# ============================================================ SLIDE 17 — SWOT
s = content_slide("SWOT", "SWOT 分析", 17)
chip(s, Inches(0.55), Inches(1.5), Inches(6.05), Inches(2.45), "S 优势 Strengths",
     ["• 豆包模型(120万亿 token/全球前三)",
      "• 火山引擎(调用份额 59.2%)",
      "• 抖音/飞书流量分发独一无二",
      "• 开源生态(Studio 约 2.1万 Star)",
      "• 易用 + 多模型开放"], GREEN)
chip(s, Inches(6.75), Inches(1.5), Inches(6.05), Inches(2.45), "W 劣势 Weaknesses",
     ["• 工具型变现天花板、低粘性",
      "• 两年内定位多次摇摆",
      "• 与竞品功能高度同质化",
      "• 深度绑定字节生态(迁移成本)",
      "• 国内/海外版本割裂"], AMBER)
chip(s, Inches(0.55), Inches(4.1), Inches(6.05), Inches(2.45), "O 机会 Opportunities",
     ["• MCP/A2A 标准化(扣子已支持 MCP)",
      "• 2026 被视为 Agent 规模化落地元年",
      "• 职场 AI / Vibe Coding 新赛道",
      "• 火山引擎主战场押注 Agent",
      "• 企业 To B 私有化需求(HiAgent)"], BLUE)
chip(s, Inches(6.75), Inches(4.1), Inches(6.05), Inches(2.45), "T 威胁 Threats",
     ["• 母公司利润波动加大算力投入压力",
      "• 腾讯 ADP3.0/阿里百炼/Dify 夹击",
      "• 模型进步“稀释工作流价值”悖论",
      "• 监管/合规(出海面临欧盟 AI 法案)",
      "• 同质化致价格战风险"], RED)

# ============================================================ SLIDE 18 — RISKS
s = content_slide("RISKS & CHALLENGES", "风险与挑战（诚实披露）", 18)
bullets(s, [
 {"lead":"母公司利润压力 — ", "t":"有报道称字节 2025 净利润同比下滑超 70%(主因 AI 算力/芯片采购猛增)；官方回应称系国际会计准则口径(含优先股/期权成本)，不反映运营实质。算力“吞金”是宏观风险。", "color":AMBER, "tag":"中-强(含官方部分否认)"},
 {"lead":"商业化天花板与低粘性 — ", "t":"分析指工具型产品订阅付费逻辑挑战大、LTV 天花板低、低频低粘性。", "color":AMBER, "tag":"中(观点)"},
 {"lead":"“模型进步稀释工作流价值”悖论 — ", "t":"模型越弱、工作流越有意义；模型不断进步，手工编排的价值被稀释。", "color":AMBER, "tag":"中(观点)"},
 {"lead":"定位摇摆 — ", "t":"两年内 Bot 工厂 → 通用 Agent → 职场 AI + Vibe Coding 多次重定位，被报道描述为“走一步看一步”。", "color":AMBER, "tag":"中"},
 {"lead":"同质化 + 生态绑定双刃 — ", "t":"与腾讯/百度/阿里/Dify 在工作流+插件+知识库+商店上高度趋同，差异主要落在各自流量入口；深绑字节生态既是冷启动优势，也带来锁定与迁移成本担忧。", "color":AMBER, "tag":"中"},
], y=Inches(1.5), size=13, gap=10)

# ============================================================ SLIDE 19 — INDUSTRY TRENDS
s = content_slide("INDUSTRY TRENDS", "行业趋势 2025–2026：对扣子的影响", 19)
bullets(s, [
 {"lead":"MCP 成事实标准 — ", "t":"2025.12 MCP 捐入 Linux 基金会，多巨头共治；2026 成熟度快速追平 Docker。扣子空间/2.0 已全面支持 MCP，顺势受益。", "tag":"中-强"},
 {"lead":"多智能体进入生产 — ", "t":"2026 被视为 Multi-Agent 进生产环境、A2A/AP2 协议标准化的关键年，出现“Agent 控制平面/多 Agent 仪表盘”。", "tag":"中"},
 {"lead":"从“可用”到“规模化生产力” — ", "t":"2026 被多份研报称为 Agent 规模化落地元年，深入企业核心流程、超越 RPA。与扣子 2.0“职场 AI”定位一致。", "tag":"中"},
 {"lead":"火山引擎押注 Agent — ", "t":"豆包破 50 万亿 token 后，火山引擎宣布主战场转向 Agent 落地，与扣子/HiAgent 战略协同强化。", "tag":"中-强"},
 {"t":"含义：扣子的重心正从 C 端创作迁向“职场/企业生产力 + Vibe Coding”，踩在“Agent 规模化元年 + MCP 标准化”两大趋势上。", "lvl":1, "color":BLUE, "bullet":"➔  ", "bold":True},
], y=Inches(1.5), size=13.5, gap=11)

# ============================================================ SLIDE 20 — KEY JUDGMENTS & RECOMMENDATIONS
s = content_slide("KEY JUDGMENTS", "核心洞察与建议（分析层）", 20)
bullets(s, [
 {"lead":"洞察① 扣子的本质是字节 AI 闭环的“应用层抓手” — ", "t":"其价值不只在工具本身，而在把豆包模型、火山算力、抖音流量串成可变现闭环。脱离这条闭环看扣子会低估其战略权重。"},
 {"lead":"洞察② 开源是攻守兼备的一步 — ", "t":"对外抢 Agent 开发“事实标准”、压制 Dify；对内为企业版/云服务引流。是当前最清晰、最具杠杆的战略动作。"},
 {"lead":"洞察③ 真正的战役在“长时程 + 企业生产力” — ", "t":"竞争从“谁更易用”转向“谁能在职场/企业核心流程稳定跑通多 Agent”。2.0/2.5 的职场 AI、Agent World 即押此方向。"},
 {"lead":"建议(企业用户)— ", "t":"轻量/创作/快速验证用扣子(SaaS)；数据敏感/合规/私有化用 HiAgent 或开源自托管；多模型自由度需求高者，评估 Dify 等中立平台并行。", "color":BLUE},
 {"lead":"建议(观察/投资)— ", "t":"盯三个信号：① 企业版/专业版付费转化与营收披露；② 开源社区活跃度(Star/贡献/企业自托管)；③ 与抖音电商等字节场景的深度变现案例。", "color":BLUE},
], y=Inches(1.5), size=13, gap=10)

# ============================================================ SLIDE 21 — SOURCES & CONFIDENCE
s = content_slide("SOURCES & CONFIDENCE", "数据来源与可信度说明（节选）", 21,
                  note="完整来源见随附研究记录；以下为各维度最权威/最具价值来源")
bullets(s, [
 {"lead":"官方 / 一手(高可信)— ", "t":"github.com/coze-dev/coze-studio、coze-loop(Stars/版本 API 实测) · coze.cn 官方文档 · volcengine.com 火山引擎扣子文档/产品页 · cloudwego.io(Eino)", "size":12},
 {"lead":"发布会 / 厂商披露 — ", "t":"火山引擎 Force 大会(2024.12 / 2025.06 / 2025.12) · 扣子 2.0 发布稿(2026.01) · 豆包大模型 1.6 发布(2025.06) · 火山引擎扣子企业交流日(2025.06)", "size":12},
 {"lead":"权威媒体 / 智库 — ", "t":"极客公园《扣子两年生长真相》 · 量子位(开源拆箱/年度 AI 100) · 36氪 · InfoQ · 钛媒体 · 雷峰网 · 新浪/腾讯财经 · IDC(MaaS 份额)", "size":12},
 {"lead":"可信度处理 — ", "t":"GitHub 数字为 API 直读(强)；定价、458万月活、首日50万、300万开发者等单源/时效数字标注为“中”，建议以官网/官方稿复核。", "size":12, "color":AMBER},
 {"t":"重点提示：务必区分“豆包(C 端 App)”与“扣子(开发者平台)”数据；开发者数量随时间存在“活跃/月活/用户/开发场景”口径切换。", "lvl":1, "color":RED, "bullet":"⚠  ", "bold":True, "size":12},
 {"t":"研究方式：5 路并行检索 + 多源交叉验证 + 可信度全程分层标注。", "lvl":1, "color":GREEN, "bullet":"✔  ", "bold":True, "size":12},
], y=Inches(1.5), size=12, gap=10)

# ============================================================ SLIDE 22 — CLOSING
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, SW, SH, NAVY)
rect(s, 0, Inches(3.5), SW, Pt(3), BLUE)
rect(s, Inches(0.0), 0, Inches(0.22), SH, BLUE)
tb, tf = textbox(s, Inches(0.9), Inches(2.5), Inches(11.5), Inches(2.5))
p = tf.paragraphs[0]; style(p, "扣子 Coze · 深度洞察", 16, BLUE, bold=True, space_after=8)
p = tf.add_paragraph(); style(p, "一个平台，一条字节 AI 闭环", 34, WHITE, bold=True, space_after=10)
p = tf.add_paragraph(); style(p, "模型(豆包) → 应用(扣子) → 分发(抖音/飞书) → 算力(火山引擎)", 15, LGRAY, space_after=6)
p = tf.add_paragraph(); style(p, "综合全景洞察 · 2026 年 6 月 · 多源交叉验证 · 引用前请以官网/官方发布稿复核确切数字", 11, GRAY)

out = "/home/user/openai-chatgpt-codex/字节扣子Coze深度洞察分析.pptx"
prs.save(out)
print("saved:", out, "slides:", len(prs.slides._sldIdLst))
