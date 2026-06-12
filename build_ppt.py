# -*- coding: utf-8 -*-
"""Generate a 2-slide, high-density technical insight PPT for GenericAgent."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

NAVY   = RGBColor(0x0B, 0x2A, 0x42)
BLUE   = RGBColor(0x12, 0x5E, 0x8C)
TEAL   = RGBColor(0x00, 0xB3, 0xA4)
ORANGE = RGBColor(0xFF, 0x7A, 0x33)
LIGHT  = RGBColor(0xF3, 0xF5, 0xF8)
GREY   = RGBColor(0x5A, 0x66, 0x73)
DARK   = RGBColor(0x1A, 0x24, 0x33)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
CARDBG = RGBColor(0xED, 0xF1, 0xF6)
OPINBG = RGBColor(0x10, 0x33, 0x4E)   # opinion band bg
FONT = "Microsoft YaHei"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def set_font(run, name=FONT):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", name)


def add_rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    sp.line.fill.background(); sp.shadow.inherit = False
    return sp


def add_text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             space_after=3, line_spacing=1.0):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(2); tf.margin_top = tf.margin_bottom = Pt(1)
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space_after); p.line_spacing = line_spacing
        for (text, size, color, bold) in para:
            r = p.add_run(); r.text = text
            r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
            set_font(r)
    return tb


def title_bar(slide, kicker, title):
    add_rect(slide, 0, 0, SW, Inches(0.96), NAVY)
    add_rect(slide, 0, Inches(0.96), SW, Pt(3.5), TEAL)
    add_rect(slide, Inches(0.45), Inches(0.20), Pt(6), Inches(0.56), ORANGE)
    add_text(slide, Inches(0.66), Inches(0.11), Inches(12.2), Inches(0.8),
             [[(kicker, 10, TEAL, True)], [(title, 18.5, WHITE, True)]], space_after=1)


def panel(slide, x, y, w, h, head, lines, accent=TEAL, lh=1.02):
    """lines: list of (text, bold) — rendered as compact bullets."""
    add_rect(slide, x, y, w, h, CARDBG, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(slide, x, y, Pt(4.5), h, accent)
    runs = [[(head, 12.5, NAVY, True)]]
    for txt, bold in lines:
        runs.append([("▸ ", 9, accent, True), (txt, 9.7, DARK, bold)])
    add_text(slide, x+Inches(0.16), y+Inches(0.10), w-Inches(0.26), h-Inches(0.16),
             runs, space_after=2.5, line_spacing=lh)


def opinion(slide, head_runs, body_runs):
    y = Inches(5.78)
    add_rect(slide, 0, y, SW, SH-y, OPINBG)
    add_rect(slide, 0, y, SW, Pt(3), ORANGE)
    add_text(slide, Inches(0.5), y+Inches(0.10), Inches(12.4), Inches(0.32),
             head_runs, space_after=1)
    add_text(slide, Inches(0.5), y+Inches(0.46), Inches(12.4), Inches(1.18),
             body_runs, space_after=3, line_spacing=1.12)


# ============================= SLIDE 1 =============================
s = prs.slides.add_slide(BLANK)
title_bar(s, "ARCHITECTURE · MECHANISM　|　GenericAgent 技术洞察（1/2）",
          "极简架构：~100 行循环驱动 9 原子工具，五层记忆支撑「任务即技能」的自进化")

# left: 执行循环
panel(s, Inches(0.45), Inches(1.22), Inches(4.05), Inches(4.40), "执行循环  Agent Loop (~100 行)", [
    ("① 调 LLM（消息历史 + 工具 schema）", False),
    ("② 解析响应中的 tool calls", False),
    ("③ dispatch → do_{tool} 分发执行", False),
    ("④ 生成器 yield，exhaust 排空回收", False),
    ("⑤ 合并各工具 next_prompt 拼下一轮", False),
    ("⑥ 追加 user 消息，循环至 should_exit", False),
    ("机制：流式输出 + 生命周期 Hook", True),
    ("机制：正则后处理压缩冗长代码块", True),
    ("自进化：成功路径蒸馏为 Skill/SOP，", False),
    ("　相似任务下一次「一行直接复用」", False),
], BLUE, lh=1.18)

# mid: 9 tools
panel(s, Inches(4.62), Inches(1.22), Inches(4.05), Inches(4.40), "九大原子工具  Atomic Tools", [
    ("code_run — 执行任意 Python/PowerShell", True),
    ("file_read / file_write / file_patch", False),
    ("web_scan — 感知解析网页内容", False),
    ("web_execute_js — JS 控真实浏览器", False),
    ("ask_user — 人在环路确认", False),
    ("update_working_checkpoint — 短期记忆", False),
    ("start_long_term_update — 蒸馏持久记忆", False),
    ("核心：以 code_run「现写脚本」取代", True),
    ("　数百专用工具 → 工具极少、能力极广", True),
    ("TMWebDriver：WS 服务+Chrome 扩展注入", False),
    ("　真实会话，过 SannySoft/FingerprintJS", False),
], TEAL, lh=1.18)

# right: 5-layer memory
panel(s, Inches(8.79), Inches(1.22), Inches(4.10), Inches(4.40), "五层结晶式记忆  5-Layer Memory", [
    ("L0 元规则 — 行为约束/系统红线", True),
    ("L1 洞察索引 — 快速路由与召回入口", True),
    ("L2 全局事实 — 跨任务稳定知识", False),
    ("L3 任务技能/SOP — 可复用工作流", True),
    ("L4 会话归档 — 蒸馏长程记录", False),
    ("效果：分层「浓缩」记忆优于冗余堆叠", False),
    ("　与纯向量检索，降噪并保住可复用经验", False),
    ("规模：~3K 行种子代码 / ~30K 上下文", True),
    ("栈：Python3.11–3.12 · requests · bs4", False),
    ("　bottle · aiohttp，无 Playwright/LangChain", False),
], ORANGE, lh=1.18)

opinion(s,
    [[("▍观点 · 架构评析", 12.5, TEAL, True),
      ("　以 code_run 为「元工具」现写脚本，把工具空间从 N 个专用接口压成「图灵完备 + 9 原语」", 10.5, WHITE, False)]],
    [[("本质是把组合爆炸交给 LLM 推理而非工程预置，代价是强依赖模型代码能力与安全沙箱。五层记忆把「提示工程」升级为可持久化的", 10.5, RGBColor(0xD5,0xDF,0xE8), False),
      ("状态机", 10.5, TEAL, True),
      ("，其中 L1 索引的召回质量是系统上限的真正瓶颈。", 10.5, RGBColor(0xD5,0xDF,0xE8), False),
      ("对 openJiuwen 的借鉴：", 10.5, ORANGE, True),
      ("走「薄工具 + 厚记忆」路线，把可复用经验沉淀为 SOP 而非堆工具；优先建好记忆分层与召回索引，再谈工具扩展。", 10.5, RGBColor(0xD5,0xDF,0xE8), False)]])

# ============================= SLIDE 2 =============================
s = prs.slides.add_slide(BLANK)
title_bar(s, "VALUE · EVALUATION · INSIGHT　|　GenericAgent 技术洞察（2/2）",
          "差异化与评测：30K 上下文 + ~6× 更低 token，真实浏览器构筑反检测护城河")

panel(s, Inches(0.45), Inches(1.22), Inches(4.05), Inches(4.40), "四大差异化亮点  Differentiators", [
    ("Token 效率：~30K vs 对手 200K–1M", True),
    ("　降噪/降幻觉/降成本，约 6× 更省", False),
    ("真实浏览器：注入真实 Chrome 会话", True),
    ("　保登录态/Cookie/指纹，反检测", False),
    ("自进化：任务结晶 Skill，越用越强", True),
    ("　无需人工干预，形成个性化技能树", False),
    ("跨平台跨模型：Claude/Gemini/Kimi", True),
    ("　/MiniMax · Win/macOS/Linux", False),
    ("仓库自身由 Agent 自举完成（含 commit）", False),
], TEAL, lh=1.20)

panel(s, Inches(4.62), Inches(1.22), Inches(4.05), Inches(4.40), "五维评测 & 对比标杆  Evaluation", [
    ("① 任务完成 & Token 效率（约 6×↓）", False),
    ("② 工具使用效率：原子集 > 专用堆叠", False),
    ("③ 记忆有效性：分层浓缩 > 向量检索", False),
    ("④ 自进化：经验蒸馏为可复用 SOP", False),
    ("⑤ 网页浏览：极简下仍具开放网络韧性", False),
    ("———— 对比 Claude Code / OpenClaw ————", True),
    ("代码体量：3K 行  vs  530K 行", True),
    ("上下文：~30K  vs  200K–1M", True),
    ("Token 消耗：约 6× 更低", True),
], BLUE, lh=1.22)

panel(s, Inches(8.79), Inches(1.22), Inches(4.10), Inches(4.40), "优势 / 风险速览  SWOT", [
    ("＋ 信息密度高、成本/幻觉双低", True),
    ("＋ 自进化沉淀复用经验为个人资产", False),
    ("＋ 真实浏览器=反检测护城河", False),
    ("＋ 依赖极少、部署轻、可商用", False),
    ("－ code_run 任意执行=高权限风险", True),
    ("－ 强依赖模型推理质量", False),
    ("－ 记忆漂移/污染需治理", False),
    ("－ 评测多为自报，缺第三方基准", False),
    ("－ 自进化结果可复现性存疑", False),
    ("维护压力：65+ PR / 76 Issues", False),
], ORANGE, lh=1.22)

opinion(s,
    [[("▍观点 · 价值评析 & 对 openJiuwen 的借鉴意义", 12.5, TEAL, True)]],
    [[("30K 上下文并非省钱噱头，而是用「信息密度」换「信噪比」——更短上下文直接降低注意力稀释与幻觉，是成功率的", 10.5, RGBColor(0xD5,0xDF,0xE8), False),
      ("因而非果", 10.5, TEAL, True),
      ("；真实浏览器注入以合规/安全换反检测，是落地差异点也是最大合规风险。", 10.5, RGBColor(0xD5,0xDF,0xE8), False),
      ("借鉴：", 10.5, ORANGE, True),
      ("可移植「任务→结晶 Skill」闭环 + 分层记忆，但须配套权限沙箱、记忆审计与第三方基准——否则高权限执行与自报评测会成为规模化前的硬伤。", 10.5, RGBColor(0xD5,0xDF,0xE8), False)]])

prs.save("GenericAgent_技术洞察分析.pptx")
print("Saved 2-slide deck. Slides:", len(prs.slides._sldIdLst))
